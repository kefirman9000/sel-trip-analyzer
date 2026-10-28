"""Tests for parsing, cause analysis, report export and the updater.

Run:  python -m unittest discover -s tests -v

Real relay files: put scrubbed copies in tests/fixtures/<case-name>/ along with
an expected.json describing what the summary must say (see tests/fixtures/README.md).
Every case folder is picked up automatically by RealFileCases.
"""
import hashlib
import io
import json
import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

import sel_trip_analyzer as app  # noqa: E402

SAMPLES = os.path.join(ROOT, "sample_files")
FIXTURES = os.path.join(HERE, "fixtures")


def load(*names):
    paths = app.expand_paths([os.path.join(SAMPLES, n) for n in names])
    return app.load_files(paths) + (paths,)


class CevParsing(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        events, cls.rdb, cls.err, _ = load("test.cev")
        cls.ev = events[0]

    def test_header_fields(self):
        e = self.ev
        self.assertEqual(e.event_type, "AG T")
        self.assertEqual(e.location, "4.21")
        self.assertEqual(e.targets, "INST ZONE1 A G")
        self.assertEqual(e.date, "10/27/2026")
        self.assertTrue(e.fid.startswith("SEL-351"))
        self.assertEqual((e.relay_id, e.terminal_id), ("FEEDER 12", "APEX SUB"))

    def test_digital_sequence(self):
        e = self.ev
        self.assertAlmostEqual(e.trip_cycle, 4.25)
        self.assertAlmostEqual(e.open_cycle, 7.5)
        first = {el: cyc for cyc, el, st in e.sequence if st == "asserted"}
        self.assertLess(first["51G"], first["Z1G"])
        self.assertEqual(first["Z1G"], first["TRIP"])

    def test_magnitudes(self):
        pre, fault = self.ev.currents["IA"]
        self.assertAlmostEqual(pre, 200, delta=5)
        self.assertAlmostEqual(fault, 2500, delta=25)
        vpre, vmin = self.ev.voltages["VA(kV)"]
        self.assertAlmostEqual(vmin, 12, delta=0.5)

    def test_cause_names_trip_equation_element(self):
        lines, tripped = app.determine_cause(self.ev, self.ev.settings)
        self.assertTrue(tripped)
        text = "\n".join(lines)
        self.assertIn("Phase A-to-ground fault", text)
        self.assertIn("(TR): Z1G", text)
        self.assertIn("3.25 cycles after trip", text)


class TextParsing(unittest.TestCase):
    def test_summary_report(self):
        events, *_ = load("sum.txt")
        e = events[0]
        self.assertEqual(e.event_type, "TRIP")
        self.assertEqual(e.targets, "COMM")
        self.assertEqual(e.relay_id, "FEEDER 7")
        self.assertIn("communications-assisted", app.short_cause(e, {}))

    def test_history_rows(self):
        events, *_ = load("his.his")
        self.assertEqual([e.event_type for e in events], ["BC T", "ER"])
        self.assertEqual(events[0].location, "7.10")
        self.assertTrue(app.short_cause(events[0], {}).startswith("TRIP"))
        self.assertTrue(app.short_cause(events[1], {}).startswith("No trip"))

    def test_event_type_decoding(self):
        cases = {"AG T": ("Phase A-to-ground fault", True), "ABC": ("Three-phase fault", False),
                 "BCG T": ("Phase B-C-to-ground fault", True), "ER": ("ER", False)}
        for code, (want, trip) in cases.items():
            desc, tripped = app.interpret_event_type(code)
            self.assertIn(want, desc, code)
            self.assertEqual(tripped, trip, code)

    def test_rdb_text_fallback(self):
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "x.rdb")
            with open(p, "wb") as f:
                f.write(b"\x00\x01junk RID     =FEEDER 12\nTR      =M1P + Z1G\nLL      =10.00\n")
            s = app.parse_rdb(p)["FEEDER 12"]
            self.assertEqual((s["TR"], s["LL"]), ("M1P + Z1G", "10.00"))


class ReportAndExport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        events, rdb, err, paths = load("")
        cls.text = app.build_report(events, rdb, err, paths, "Crew found a cracked insulator.")

    def test_report_sections(self):
        for h in ("QUICK SUMMARY", "EVENT DETAILS", "ANALYST COMMENTS / NOTES", "Crew found"):
            self.assertIn(h, self.text)

    def test_blocks(self):
        blocks = app.parse_report_text(self.text)
        kinds = [b[0] for b in blocks]
        self.assertEqual(kinds[0], "title")
        qs = next(b for b in blocks if b[0] == "table" and b[1][0] == "#")
        self.assertEqual(len(qs[2]), 4)
        seq = next(b for b in blocks if b[0] == "table" and b[1][0] == "Cycle")
        self.assertIn(["4.25", "Z1G", "asserted", "distance"], seq[2])

    def test_docx_keeps_user_edits(self):
        from docx import Document
        edited = self.text.replace("QUICK SUMMARY", "QUICK SUMMARY\nEdited by hand: confirmed with dispatch.")
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "r.docx")
            app.export_report(edited, p)
            doc = Document(p)
            body = "\n".join(par.text for par in doc.paragraphs)
            cells = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn("Edited by hand: confirmed with dispatch.", body)
        self.assertIn("Quick Summary", body)
        self.assertIn("FEEDER 12", cells)

    def test_pdf(self):
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "r.pdf")
            app.export_report(self.text, p)
            with open(p, "rb") as f:
                data = f.read()
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertGreater(len(data), 3000)

    def test_garbage_text_still_exports(self):
        junk = "just some notes\n\nNOT A REPORT\n  weird : line\n    - bullet?"
        with tempfile.TemporaryDirectory() as d:
            for ext in (".docx", ".pdf", ".txt"):
                app.export_report(junk, os.path.join(d, "j" + ext))


class Updater(unittest.TestCase):
    def test_version_compare(self):
        self.assertGreater(app.version_tuple("v1.10.0"), app.version_tuple("1.9.9"))
        self.assertEqual(app.version_tuple("v1.1.0"), app.version_tuple("1.1.0"))

    def test_release_lookup(self):
        exe = b"MZ" + b"x" * 500
        rel = {"tag_name": "v99.0.0", "body": "notes", "html_url": "",
               "assets": [{"name": app.EXE_ASSET_NAME, "browser_download_url": "EXE",
                           "digest": "sha256:" + hashlib.sha256(exe).hexdigest()}]}

        class R(io.BytesIO):
            headers = {}

            def __enter__(self):
                return self

            def __exit__(self, *a):
                pass

        orig = app._http_get
        app._http_get = lambda url, **k: R(json.dumps(rel).encode() if "api.github" in url else exe)
        try:
            info = app.check_for_update()
            self.assertEqual(info["version"], "99.0.0")
            with tempfile.TemporaryDirectory() as d:
                digest = app.download_file("EXE", os.path.join(d, "x.exe"))
            self.assertEqual(digest, info["sha256"])
            rel["tag_name"] = "v" + app.VERSION
            self.assertIsNone(app.check_for_update())
        finally:
            app._http_get = orig


class RealFileCases(unittest.TestCase):
    """One sub-test per tests/fixtures/<case>/expected.json."""

    def test_cases(self):
        cases = sorted(d for d in os.listdir(FIXTURES)
                       if os.path.isfile(os.path.join(FIXTURES, d, "expected.json"))) \
            if os.path.isdir(FIXTURES) else []
        if not cases:
            self.skipTest("no real-file fixtures yet")
        for case in cases:
            folder = os.path.join(FIXTURES, case)
            with self.subTest(case=case), open(os.path.join(folder, "expected.json")) as f:
                exp = json.load(f)
                paths = app.expand_paths([folder])
                events, rdb, err = app.load_files(paths)
                self.assertFalse(err, err)
                evs = [e for e in events if e.kind in ("event", "history")]
                if "event_count" in exp:
                    self.assertEqual(len(evs), exp["event_count"])
                report = app.build_report(events, rdb, err, paths)
                for want in exp.get("report_contains", []):
                    self.assertIn(want, report)
                for i, ee in enumerate(exp.get("events", [])):
                    e = next((x for x in evs if x.source == ee["file"]), None) if "file" in ee else evs[i]
                    self.assertIsNotNone(e, ee)
                    for key in ("event_type", "location", "targets", "relay_id"):
                        if key in ee:
                            self.assertEqual(getattr(e, key), ee[key], f"{case} {key}")
                    cause = "\n".join(app.determine_cause(e, {**e.settings})[0])
                    for want in ee.get("cause_contains", []):
                        self.assertIn(want, cause, f"{case}: cause should mention {want!r}")


if __name__ == "__main__":
    unittest.main()
