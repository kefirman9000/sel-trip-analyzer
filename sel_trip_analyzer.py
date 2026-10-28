#!/usr/bin/env python3
"""
SEL Trip Event Analyzer
-----------------------
Drag-and-drop Windows tool that reads SEL relay event files and AcSELerator
QuickSet .RDB settings databases, then writes an editable, copy-and-paste
summary of what caused each trip.

Supported inputs
  .CEV            Compressed event reports (CEV / CEV R)
  .EVE .TXT .SUM  Text event reports / event summaries (EVE, SUM, CHI, etc.)
  .HIS            Event history (HIS command output)
  .SER            Sequential Events Recorder output
  .HDR            COMTRADE header (SEL writes the event summary here)
  .TXT            Saved terminal captures of EVE / SUM output
  .RDB            QuickSet relay database (settings context: TR, LL, CTR, ...)
  Folders         Dropped folders are scanned for the types above
"""

import csv
import datetime
import math
import os
import re
import sys
import zipfile
from dataclasses import dataclass, field

try:
    import olefile  # reads the OLE container used by QuickSet .rdb files
except ImportError:  # pragma: no cover
    olefile = None

APP_NAME = "SEL Trip Event Analyzer"
VERSION = "1.2.0"
# GitHub "owner/repo" that publishes releases. Each release must attach
# SEL_Trip_Analyzer.exe (the GitHub Actions workflow in this repo does that).
GITHUB_REPO = "kefirman9000/sel-trip-analyzer"
EXE_ASSET_NAME = "SEL_Trip_Analyzer.exe"


# --------------------------------------------------------------------------
# Self-updater (GitHub Releases)
# --------------------------------------------------------------------------
def version_tuple(v):
    nums = re.findall(r"\d+", str(v))
    return tuple(int(n) for n in nums[:4]) or (0,)


def _http_get(url, accept="application/vnd.github+json", timeout=20):
    import urllib.request
    req = urllib.request.Request(url, headers={
        "Accept": accept, "User-Agent": f"SEL-Trip-Analyzer/{VERSION}"})
    return urllib.request.urlopen(req, timeout=timeout)


def check_for_update():
    """Return None if up to date, else dict(version, notes, url, sha256, page, source_url)."""
    import json
    with _http_get(f"https://api.github.com/repos/{GITHUB_REPO}/releases/latest") as r:
        rel = json.load(r)
    tag = rel.get("tag_name", "")
    if version_tuple(tag) <= version_tuple(VERSION):
        return None
    asset = next((a for a in rel.get("assets", []) if a.get("name") == EXE_ASSET_NAME), None)
    digest = (asset or {}).get("digest") or ""
    return {
        "version": tag.lstrip("vV"),
        "tag": tag,
        "notes": (rel.get("body") or "").strip(),
        "url": asset["browser_download_url"] if asset else None,
        "sha256": digest.split(":", 1)[1].lower() if digest.startswith("sha256:") else None,
        "page": rel.get("html_url", ""),
        "source_url": f"https://raw.githubusercontent.com/{GITHUB_REPO}/{tag}/sel_trip_analyzer.py",
    }


def download_file(url, dest, progress=None):
    import hashlib
    h = hashlib.sha256()
    with _http_get(url, accept="application/octet-stream", timeout=60) as r, open(dest, "wb") as f:
        total = int(r.headers.get("Content-Length") or 0)
        done = 0
        while True:
            chunk = r.read(65536)
            if not chunk:
                break
            f.write(chunk)
            h.update(chunk)
            done += len(chunk)
            if progress:
                progress(done, total)
    return h.hexdigest()


def install_update(info, progress=None):
    """Download and stage the update. Returns a callable that finishes the
    install and restarts the app (call it after closing the window)."""
    import subprocess
    import tempfile

    if getattr(sys, "frozen", False):  # running as the built .exe
        if not info["url"]:
            raise RuntimeError("This release has no .exe attached yet - try again in a few minutes.")
        exe = sys.executable
        folder = os.path.dirname(exe)
        if not os.access(folder, os.W_OK):
            raise RuntimeError(f"No permission to write to {folder}.\n"
                               "Move SEL_Trip_Analyzer.exe to a folder you own (e.g. Desktop) and retry.")
        new_exe = exe + ".new"
        digest = download_file(info["url"], new_exe, progress)
        if info["sha256"] and digest != info["sha256"]:
            os.remove(new_exe)
            raise RuntimeError("Downloaded file failed its checksum - update cancelled.")
        bat = os.path.join(tempfile.gettempdir(), "sel_trip_analyzer_update.bat")
        with open(bat, "w") as f:
            f.write(
                "@echo off\r\n"
                ":wait\r\n"
                f'tasklist /FI "PID eq {os.getpid()}" 2>nul | find "{os.getpid()}" >nul && '
                "(timeout /t 1 /nobreak >nul & goto wait)\r\n"
                f'move /Y "{new_exe}" "{exe}" >nul\r\n'
                f'start "" "{exe}"\r\n'
                'del "%~f0"\r\n')

        def finish():
            flags = 0x08000000 | 0x00000008  # CREATE_NO_WINDOW | DETACHED_PROCESS
            subprocess.Popen(["cmd", "/c", bat], creationflags=flags, close_fds=True)
        return finish

    # running from source: replace this .py file
    script = os.path.abspath(__file__)
    tmp = script + ".new"
    download_file(info["source_url"], tmp, progress)
    with open(tmp, encoding="utf-8") as f:
        compile(f.read(), tmp, "exec")  # make sure it is valid Python before swapping
    os.replace(script, script + ".bak")
    os.replace(tmp, script)

    def finish():
        subprocess.Popen([sys.executable, script])
    return finish

EVENT_EXTS = {".cev", ".eve", ".txt", ".sum", ".his", ".ser", ".hdr", ".log", ".chi"}
SETTINGS_EXTS = {".rdb"}
ALL_EXTS = EVENT_EXTS | SETTINGS_EXTS


# --------------------------------------------------------------------------
# Data model
# --------------------------------------------------------------------------
@dataclass
class EventRecord:
    source: str
    kind: str = "event"           # event | history | ser
    relay_id: str = ""
    terminal_id: str = ""
    fid: str = ""
    date: str = ""
    time: str = ""
    event_type: str = ""
    location: str = ""
    shot: str = ""
    freq: str = ""
    targets: str = ""
    currents: dict = field(default_factory=dict)        # name -> (prefault, fault max)
    voltages: dict = field(default_factory=dict)        # name -> (prefault, fault min)
    sequence: list = field(default_factory=list)        # (cycle, element, 'asserted'/'pre')
    trip_cycle: float = None
    open_cycle: float = None
    settings: dict = field(default_factory=dict)
    ser_lines: list = field(default_factory=list)
    extra: list = field(default_factory=list)
    dup_of: int = None                                  # history row that repeats event #n
    trigger_cycle: float = None
    prefault_short: bool = False                        # trigger too early for a clean prefault
    notes: list = field(default_factory=list)           # extra cause lines (e.g. duplicates)


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------
def read_text(path):
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def is_num(s):
    try:
        float(str(s).strip())
        return True
    except ValueError:
        return False


def fnum(s, default=None):
    try:
        return float(str(s).strip())
    except (ValueError, TypeError):
        return default


SETTING_RE = re.compile(
    r"\b([A-Z][A-Z0-9_]{0,11})\s*:?=\s*(.*?)(?=\s{2,}[A-Z][A-Z0-9_]{0,11}\s*:?=|\"|$)"
)
USEFUL_SETTINGS = {
    "RID", "TID", "SID", "LL", "CTR", "CTRN", "PTR", "PTRS", "Z1MAG", "Z1ANG", "Z0MAG",
    "TR", "TRCOMM", "TRSOTF", "DTT", "ULTR", "ER", "79RI", "52A", "52AA1", "BFI",
    "Z1P", "Z2P", "Z3P", "Z1MG", "Z2MG", "Z1G", "Z2D", "Z2PD", "Z2GD", "Z3D",
    "50P1P", "50G1P", "50Q1P", "51PP", "51PC", "51PTD", "51GP", "51GC", "51GTD",
    "51QP", "E21P", "E21MG", "E21XG", "E50P", "E50G", "E51S", "E32", "ECOMM", "E81",
    "TOP", "TPA", "TPB", "TPC",
}


def extract_settings(text):
    out = {}
    for line in text.splitlines():
        if "=" not in line:
            continue
        for m in SETTING_RE.finditer(line):
            key, val = m.group(1), m.group(2).strip().strip(",").strip()
            if not val or len(val) > 400:
                continue
            if key in USEFUL_SETTINGS or key.startswith(("TR", "E", "Z", "50", "51", "67", "87")):
                out.setdefault(key, val)
    return out


# --------------------------------------------------------------------------
# Interpretation tables
# --------------------------------------------------------------------------
FAULT_TYPES = {
    "AG": "Phase A-to-ground fault", "BG": "Phase B-to-ground fault", "CG": "Phase C-to-ground fault",
    "AB": "Phase A-B (phase-to-phase) fault", "BC": "Phase B-C (phase-to-phase) fault",
    "CA": "Phase C-A (phase-to-phase) fault",
    "ABG": "Phase A-B-to-ground fault", "BCG": "Phase B-C-to-ground fault", "CAG": "Phase C-A-to-ground fault",
    "ABC": "Three-phase fault", "ABCG": "Three-phase-to-ground fault",
}
NON_FAULT_TYPES = {
    "TRIP": "Trip with no fault type identified (non-fault trip: transfer/remote/manual trip, "
            "external input, or fault type undetermined)",
    "ER": "Event triggered by the ER (event report trigger) equation - no trip",
    "TRIG": "Manually triggered event (TRI command) - no trip",
    "PULSE": "Output pulsed by command (PUL)",
    "OPEN": "Breaker open command", "CLOSE": "Breaker close command",
    "BCW": "Breaker contact wear alarm", "LOP": "Loss of potential",
    "OST": "Out-of-step trip", "OSB": "Out-of-step blocking",
    "SOTF": "Switch-onto-fault",
}

TARGET_WORDS = [
    (r"\bINST\b", "instantaneous element"),
    (r"\bCOMM\b", "communications-assisted scheme (POTT/DCB/DCUB/DTT)"),
    (r"\bSOTF\b", "switch-onto-fault logic"),
    (r"\bTIME\b", "time-delayed element"),
    (r"\bZ(?:ONE)?\s*1\b|\bZONE1\b|\bZ1\b", "Zone 1 distance"),
    (r"\bZ(?:ONE)?\s*2\b|\bZONE2\b|\bZ2\b", "Zone 2 distance"),
    (r"\bZ(?:ONE)?\s*3\b|\bZONE3\b|\bZ3\b", "Zone 3 distance"),
    (r"\bZ(?:ONE)?\s*4\b|\bZONE4\b|\bZ4\b", "Zone 4 distance"),
    (r"\b87\w*\b", "differential (87)"),
    (r"\b51\w*\b", "time-overcurrent (51)"),
    (r"\b50\w*\b", "instantaneous overcurrent (50)"),
    (r"\b67\w*\b", "directional overcurrent (67)"),
    (r"\b81\w*\b|\bFREQ\b", "frequency (81)"),
    (r"\b27\w*\b", "undervoltage (27)"),
    (r"\b59\w*\b", "overvoltage (59)"),
    (r"\bBF\w*\b|\b50BF\b", "breaker failure"),
    (r"\bCY\b", "recloser cycling (reclose in progress)"),
    (r"\bRS\b", "recloser reset (ready)"),
    (r"\b79\w*\b|\bRCLS\b", "reclosing"),
    (r"\bLO\b", "recloser locked out"),
]

ELEMENT_RULES = [
    (r"^(Z|M)\d+(P|G|MG|XG)?T$|^Z\d+(P|G)?D$", "distance (time-delayed)"),
    (r"^(Z|M)\d+(P|G|MG|XG|AB|BC|CA|AG|BG|CG)?$|^MAB\d|^MBC\d|^MCA\d|^MAG\d|^MBG\d|^MCG\d|^XAG\d|^XBG\d|^XCG\d", "distance"),
    (r"^87", "differential"),
    (r"^51\w*T$", "time-overcurrent trip"),
    (r"^51", "time-overcurrent pickup"),
    (r"^67", "directional overcurrent"),
    (r"^50BF|^BFT|^FBF", "breaker failure"),
    (r"^50", "instantaneous overcurrent"),
    (r"^32|^F32|^R32", "directional"),
    (r"^81", "frequency"),
    (r"^27", "undervoltage"),
    (r"^59", "overvoltage"),
    (r"^24", "volts/hertz"),
    (r"^46", "negative-sequence"),
    (r"^49", "thermal"),
    (r"^OST|^OSB", "out-of-step"),
    (r"^SOTF|^SOTFT", "switch-onto-fault"),
    (r"^(DTT|DTA|DTB|DTC|PTRX|PT\d?|RX|KEY|TRCOMM|ECTT|DCB|POTT|Z2PGS|RB\d|ROK|RBADA?)", "communications / transfer trip"),
    (r"^IN\d+$", "contact input"),
    (r"^(TRIP|TR|TRIP\d|3PT|TPA\d?|TPB\d?|TPC\d?|TRP\w*)$", "TRIP"),
    (r"^OUT\d+$", "output contact"),
    (r"^(52A\w*|52B\w*)$", "breaker status"),
    (r"^79", "reclosing"),
    (r"^LOP|^ILOP", "loss of potential"),
    (r"^(SV|PSV|PLT|PCT|LT|LV|PCV)\d+", "SELogic variable"),
]

TRIP_NAMES = {"TRIP", "TR", "TRIP1", "3PT", "TPA", "TPB", "TPC", "TRIP3", "TRP", "TPA1", "TRIPA"}
BREAKER_NAMES = ("52A", "52AA1", "52AB1", "52AC1", "52A1", "IN101")


def element_class(name):
    n = name.upper()
    for pat, cls in ELEMENT_RULES:
        if re.match(pat, n):
            return cls
    return ""


def interpret_event_type(et):
    et = (et or "").strip().upper()
    if not et:
        return "", False
    tripped = bool(re.search(r"\bT$|\bTRIP\b", et)) or et.endswith(" T")
    core = re.sub(r"\s+T$", "", et).strip(" *")
    core = re.sub(r"\s+", "", core)
    if core in FAULT_TYPES:
        desc = FAULT_TYPES[core]
    elif core in NON_FAULT_TYPES:
        desc = NON_FAULT_TYPES[core]
        if core == "TRIP":
            tripped = True
    else:
        desc = f"Event type '{et}'"
    return desc, tripped


def interpret_targets(t):
    if not t:
        return []
    found = []
    for pat, desc in TARGET_WORDS:
        if re.search(pat, t.upper()) and desc not in found:
            found.append(desc)
    return found


def fmt_loc(loc, settings):
    if not loc:
        return ""
    if "$" in loc:
        return f"{loc} (not calculated - relay could not determine distance)"
    v = fnum(loc)
    ll = fnum(settings.get("LL", "")) if settings else None
    if v is not None and v < 0:
        return f"{loc} (negative: fault was likely behind the relay)"
    if v is not None and ll:
        return f"{loc} line-length units (~{100 * v / ll:.0f}% of LL={settings['LL']})"
    return f"{loc} (in the relay's line-length units)"


# --------------------------------------------------------------------------
# CEV (compressed event report) parser
# --------------------------------------------------------------------------
CHECKSUM_RE = re.compile(r',\s*"[0-9A-Fa-f]{4}"\s*$')


def parse_cev(text, source):
    rows = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        has_ck = bool(CHECKSUM_RE.search(line))
        try:
            fields = next(csv.reader([line]))
        except Exception:
            continue
        fields = [f.strip() for f in fields]
        if has_ck and len(fields) > 1:
            fields = fields[:-1]
        rows.append(fields)

    rec = EventRecord(source=source)
    kv = {}
    table_hdr, table = None, []
    i = 0
    while i < len(rows):
        r = rows[i]
        if r and r[0].upper() == "FID":
            rec.fid = r[1].replace("FID=", "") if len(r) > 1 else ""
            i += 1
            continue
        if len(r) >= 2 and not any(is_num(x) for x in r):
            block = []
            j = i + 1
            while j < len(rows) and len(rows[j]) == len(r) and any(is_num(x) for x in rows[j]):
                block.append(rows[j])
                j += 1
            if len(block) == 1:
                for k, v in zip(r, block[0]):
                    kv[k.upper()] = v
            elif len(block) > len(table):
                table_hdr, table = r, block
            i = j if block else i + 1
            continue
        i += 1

    rec.event_type = kv.get("EVENT", "")
    rec.location = kv.get("LOCATION", kv.get("LOC", ""))
    rec.shot = kv.get("SHOT", kv.get("SHOT_1P", ""))
    rec.freq = kv.get("FREQ", "")
    rec.targets = kv.get("TARGETS", "")
    try:
        y, mo, d = int(kv["YEAR"]), int(kv["MONTH"]), int(kv["DAY"])
        rec.date = f"{mo:02d}/{d:02d}/{y}"
        rec.time = f"{int(kv['HOUR']):02d}:{int(kv['MIN']):02d}:{int(kv['SEC']):02d}.{int(kv.get('MSEC', 0)):03d}"
    except (KeyError, ValueError):
        pass

    rec.settings = extract_settings(text)
    rec.relay_id = rec.settings.get("RID", "")
    rec.terminal_id = rec.settings.get("TID", "")

    if table_hdr and table:
        analyze_table(rec, table_hdr, table, fnum(kv.get("SAM/CYC_A"), 4) or 4)
    return [rec]


def analyze_table(rec, hdr, table, sam):
    sam = int(sam) if sam and sam > 0 else 4
    lag = max(1, round(sam / 4))
    trig = next((k for k, row in enumerate(table) if ">" in row), None)
    start_fault = trig if trig is not None else 0
    rec.trigger_cycle = start_fault / sam if trig is not None else None
    rec.prefault_short = trig is not None and trig < sam + lag
    analog = []

    # Analog channels
    for c, name in enumerate(hdr):
        n = name.upper().split("(")[0].strip()
        if " " in name or not n:
            continue
        vals = [fnum(row[c]) for row in table]
        if any(v is None for v in vals):
            continue
        mags = [math.hypot(vals[k], vals[k - lag]) / math.sqrt(2) for k in range(lag, len(vals))]
        if not mags:
            continue
        # prefault = the cycle just before the trigger (falls back to the first cycle)
        pre_end = max(sam, start_fault - lag) if trig is not None else sam
        pre_win = mags[max(0, pre_end - sam):pre_end] or mags[:sam]
        pre = sum(pre_win) / len(pre_win)
        analog.append((name, n, pre, mags))

    # Digital channels (hex-packed relay word bits) - read before the analog
    # summary so the fault window can stop when the breaker opens
    first_on, pre_on = {}, set()
    for c, name in enumerate(hdr):
        if " " not in name.strip():
            continue
        names = name.split()
        for k, row in enumerate(table):
            h = row[c].strip() if c < len(row) else ""
            if not h or not re.fullmatch(r"[0-9A-Fa-f]+", h):
                continue
            bits = "".join(f"{int(ch, 16):04b}" for ch in h)
            for b, el in zip(bits, names):
                if el == "*" or b != "1":
                    continue
                if k == 0:
                    pre_on.add(el)
                elif el not in first_on and el not in pre_on:
                    first_on[el] = k
        # breaker opening: 52A deasserting
        for k, row in enumerate(table):
            h = row[c].strip() if c < len(row) else ""
            if not re.fullmatch(r"[0-9A-Fa-f]+", h or "-"):
                continue
            bits = "".join(f"{int(ch, 16):04b}" for ch in h)
            state = dict(zip(names, bits))
            for bn in BREAKER_NAMES[:5]:
                if bn in pre_on and state.get(bn) == "0" and rec.open_cycle is None:
                    rec.open_cycle = k / sam

    # fault window: trigger -> breaker open (if seen) else end of record
    end = int(rec.open_cycle * sam) if rec.open_cycle is not None else len(table)
    for name, n, pre, mags in analog:
        win = mags[max(0, start_fault - lag):max(start_fault - lag + 1, end - lag)] or mags
        if re.match(r"^(I[ABCNGRSPQ]|IA|IB|IC|IN|IG|3I0|3I2|I1|I2|IP)\w*$", n):
            rec.currents[name] = (pre, max(win))
        elif re.match(r"^V(A|B|C|AB|BC|CA|S|P|1|2)\w*$", n):
            rec.voltages[name] = (pre, min(win))

    for el, k in sorted(first_on.items(), key=lambda x: x[1]):
        rec.sequence.append((k / sam, el, "asserted"))
    for el in sorted(pre_on):
        if element_class(el) and element_class(el) not in ("breaker status", "SELogic variable"):
            rec.sequence.append((0.0, el, "pre"))
    trips = [k for el, k in first_on.items() if el.upper() in TRIP_NAMES]
    if trips:
        rec.trip_cycle = min(trips) / sam


# --------------------------------------------------------------------------
# Text report parsers (EVE / SUM / HIS / SER / HDR)
# --------------------------------------------------------------------------
EVENT_LINE_RE = re.compile(r"^\s*Event\s*(?:Type)?\s*:\s*(.+?)(?:\s{2,}|$)", re.I | re.M)


def grab(pattern, text, flags=re.I | re.M):
    m = re.search(pattern, text, flags)
    return m.group(1).strip() if m else ""


def parse_text(text, source):
    if '"FID"' in text or re.search(r'^"MONTH","DAY"', text, re.M):
        return parse_cev(text, source)

    recs = []
    settings = extract_settings(text)

    # Event history (HIS)
    if re.search(r"\bEVENT\b.*\bLOCAT", text, re.I):
        recs.extend(parse_his(text, source, settings))

    # SER
    ser = re.findall(
        r"^\s*\d*\s*(\d{1,2}/\d{1,2}/\d{2,4})\s+(\d{1,2}:\d{2}:\d{2}(?:\.\d+)?)\s+(\S+)\s+(Asserted|Deasserted)\b",
        text, re.I | re.M)
    if ser:
        r = EventRecord(source=source, kind="ser")
        r.ser_lines = [f"{d} {t}  {el:<12} {st}" for d, t, el, st in ser]
        recs.append(r)

    # Event reports / summaries
    matches = list(EVENT_LINE_RE.finditer(text))
    header_re = re.compile(r"^[ \t]*(\S.*?)\s{2,}Date\s*:\s*(\S+)\s+Time\s*:\s*(\S+)[^\n]*\n((?:[ \t]*\n)*)([^\n]*)", re.M)
    for idx, m in enumerate(matches):
        prev_end = matches[idx - 1].end() if idx else 0
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
        back = text[prev_end:m.start()]      # report header (relay, date/time, FID)
        chunk = text[m.start():end]          # this event's own lines
        r = EventRecord(source=source)
        heads = list(header_re.finditer(back)) or list(header_re.finditer(chunk))
        if heads:
            h = heads[-1]
            r.date, r.time = h.group(2), h.group(3)
            if not settings.get("RID"):
                r.relay_id = h.group(1).strip()
                second = h.group(5).strip()
                if second and not re.match(r"^(Event|FID|#|=>|\[)", second, re.I):
                    r.terminal_id = re.split(r"\s{2,}", second)[0].strip()
        r.event_type = re.split(r"\s{2,}|Location|Shot|Freq", m.group(1), flags=re.I)[0].strip()
        r.location = grab(r"(?:Fault\s+)?Location\s*:\s*([-$\d.]+)", chunk)
        r.shot = grab(r"Shot\s*:\s*(\S+)", chunk)
        r.freq = grab(r"Freq(?:uency)?\s*:\s*([\d.]+)", chunk)
        r.targets = grab(r"Targets?\s*:\s*(.+)$", chunk)
        if not r.date:
            r.date = grab(r"Date\s*:\s*(\S+)", chunk)
            r.time = grab(r"\bTime\s*:\s*(\d\S*)", chunk)
        if settings.get("RID"):
            r.relay_id, r.terminal_id = settings.get("RID", ""), settings.get("TID", "")
        r.fid = grab(r"FID\s*=\s*(\S+)", back) or grab(r"FID\s*=\s*(\S+)", chunk)
        for lbl in ("Trip Time", "Event Time", r"Breaker\s*\d*\s*\w*", "Fault Type", "Group", "Recloser"):
            v = grab(rf"^\s*({lbl}\s*:.+)$", chunk)
            if v:
                r.extra.append(re.sub(r"\s+", " ", v))
        cur = re.search(r"Currents?\s*\(([^)]*)\)\s*,?\s*([A-Z]+)\s*:\s*([-\d.\s]+)", chunk, re.I)
        if cur:
            labels, nums = cur.group(2), cur.group(3).split()
            for ch, v in zip(labels, nums):
                fv = fnum(v)
                if fv is not None:
                    rec_name = "3I2" if ch.upper() == "Q" else "I" + ch.upper()
                    r.currents[rec_name] = (None, fv)
        r.settings = settings
        recs.append(r)
    return recs


def parse_his(text, source, settings):
    recs = []
    head = re.match(r"\s*(.+?)\s{2,}Date\s*:", text)
    rid = settings.get("RID") or (head.group(1).strip() if head else "")
    for line in text.splitlines():
        m = re.match(r"^\s*(\d+)\s+(\d{1,2}/\d{1,2}/\d{2,4})\s+(\d{1,2}:\d{2}:\d{2}(?:\.\d+)?)\s+(.*)$", line)
        if not m:
            continue
        toks = m.group(4).split()
        et = []
        while toks and not (is_num(toks[0]) or "$" in toks[0]):
            et.append(toks.pop(0))
        r = EventRecord(source=source, kind="history")
        r.date, r.time, r.event_type = m.group(2), m.group(3), " ".join(et)
        nums = []
        while toks and (is_num(toks[0]) or "$" in toks[0]):
            nums.append(toks.pop(0))
        if nums:
            r.location = nums[0]
        if len(nums) > 1:
            r.currents["Fault current"] = (None, fnum(nums[1], 0))
        if len(nums) > 2:
            r.freq = nums[2]
        if len(nums) > 4:
            r.shot = nums[4]
        r.targets = " ".join(toks)
        r.settings = settings
        r.relay_id = rid
        recs.append(r)
    return recs


# --------------------------------------------------------------------------
# RDB (QuickSet relay database) parser
# --------------------------------------------------------------------------
def parse_rdb(path):
    """Return {relay_label: settings_dict}. Best effort across RDB variants."""
    texts = []
    with open(path, "rb") as f:
        raw = f.read()
    if olefile is not None and olefile.isOleFile(path):
        try:
            ole = olefile.OleFileIO(path)
            for entry in ole.listdir(streams=True, storages=False):
                try:
                    data = ole.openstream(entry).read()
                except Exception:
                    continue
                texts.append(("/".join(entry), decode_blob(data)))
            ole.close()
        except Exception:
            pass
    if not texts and zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                try:
                    texts.append((n, decode_blob(z.read(n))))
                except Exception:
                    continue
    method = "ole" if texts else None
    if not texts:
        texts.append(("raw", decode_blob(raw)))
        method = "scan"
    method = method or "zip"

    combined = {}
    per_relay = {}
    for name, t in texts:
        s = extract_settings(t)
        if not s:
            continue
        top = name.split("/")[0]
        per_relay.setdefault(top, {}).update({k: v for k, v in s.items() if k not in per_relay.get(top, {})})
        for k, v in s.items():
            combined.setdefault(k, v)
    if len(per_relay) <= 1:
        if not combined:
            return {}
        combined["__method__"] = method
        return {combined.get("RID") or os.path.basename(path): combined}
    out = {}
    for top, s in per_relay.items():
        s["__method__"] = method
        out[s.get("RID") or top] = s
    return out


def decode_blob(data):
    if data[:2] in (b"\xff\xfe", b"\xfe\xff") or (len(data) > 4 and data[1:2] == b"\x00" and data[3:4] == b"\x00"):
        try:
            return data.decode("utf-16", errors="ignore")
        except Exception:
            pass
    t = data.decode("latin-1", errors="ignore")
    return re.sub(r"[^\x09\x0a\x0d\x20-\x7e]+", "\n", t)


# --------------------------------------------------------------------------
# Loading and report building
# --------------------------------------------------------------------------
def expand_paths(paths):
    out = []
    for p in paths:
        if os.path.isdir(p):
            for root, _, files in os.walk(p):
                for fn in sorted(files):
                    if fn.upper().startswith("README"):
                        continue
                    if os.path.splitext(fn)[1].lower() in ALL_EXTS:
                        out.append(os.path.join(root, fn))
        elif os.path.isfile(p):
            out.append(p)
    seen, uniq = set(), []
    for p in out:
        k = os.path.normcase(os.path.abspath(p))
        if k not in seen:
            seen.add(k)
            uniq.append(p)
    return uniq


def load_files(paths):
    events, rdb, errors = [], {}, []
    for p in paths:
        ext = os.path.splitext(p)[1].lower()
        name = os.path.basename(p)
        try:
            if ext in SETTINGS_EXTS:
                rdb.update(parse_rdb(p))
            elif ext == ".cev":
                events.extend(parse_cev(read_text(p), name))
            else:
                events.extend(parse_text(read_text(p), name))
        except Exception as e:  # keep going on bad files
            errors.append(f"{name}: {e}")
    return events, rdb, errors


def match_settings(rec, rdb):
    """Pick the RDB settings for an event.

    Returns (settings, label, how) where how is:
      rid      - matched on relay ID
      assumed  - event has no relay ID and exactly one RDB is loaded
      mismatch - RDBs loaded but none match this relay (nothing applied)
      none     - no RDBs loaded
    """
    if not rdb:
        return {}, "", "none"
    rid = (rec.relay_id or "").strip().upper()
    for label, s in rdb.items():
        if rid and rid in (str(label).strip().upper(), str(s.get("RID", "")).strip().upper()):
            return s, label, "rid"
    if not rid and len(rdb) == 1:
        label = next(iter(rdb))
        return rdb[label], label, "assumed"
    return {}, "", "mismatch"


def event_core(et):
    return re.sub(r"\s+T$", "", (et or "").strip().upper()).strip(" *").replace(" ", "")


def event_phases(et):
    core = event_core(et)
    return set(core) if core in FAULT_TYPES else None


def target_phases(targets):
    toks = set((targets or "").upper().split())
    ph = {p for p in "ABC" if p in toks}
    if toks & {"G", "N"}:
        ph.add("G")
    return ph


def phase_currents(rec):
    """Split recorded currents into per-phase and ground (pre, fault) pairs."""
    ph, grounds = {}, {}
    for name, (pre, post) in rec.currents.items():
        n = name.upper().split("(")[0].strip()
        m = re.match(r"^I([ABC])(?:W\d+|_PRI|_SEC)?$", n)
        if m:
            ph[m.group(1)] = (pre, post)
        else:
            g = re.match(r"^(IG|3I0|IN)", n)
            if g:
                grounds.setdefault(g.group(1), (pre, post))
    # residual ground (IG / 3I0) reflects the fault; IN is a separate neutral CT that may read ~0
    gnd = grounds.get("IG") or grounds.get("3I0") or grounds.get("IN")
    return ph, gnd


def equation_hits(equation, elements):
    """Elements (in order) that appear un-negated in a SELogic equation."""
    eq = re.sub(r"(\bNOT\s+|!)\(?\s*[A-Z0-9_]+", " ", equation.upper())
    tokens = set(re.findall(r"[A-Z0-9_]+", eq))
    return list(dict.fromkeys(el for el in elements if el.upper() in tokens))


NON_PROTECTION = ("", "TRIP", "output contact", "breaker status", "SELogic variable", "reclosing")
TRIP_EQUATIONS = (("TR", "trip equation"), ("TRCOMM", "communications-assisted trip equation"),
                  ("TRSOTF", "switch-onto-fault trip equation"), ("DTT", "direct transfer trip equation"))


def analyze_event(rec, settings, how="none", label=""):
    """Work out the probable cause and everything a person should double-check.

    Returns dict with: tripped, desc, cause (lines), checks (problems that lower
    confidence), gaps (missing data), confidence, basis, primary (elements).
    """
    desc, tripped_by_type = interpret_event_type(rec.event_type)
    core = event_core(rec.event_type)
    is_fault = core in FAULT_TYPES
    tc = rec.trip_cycle
    tripped = tripped_by_type or tc is not None
    tdesc = interpret_targets(rec.targets)
    cause, checks, gaps = [], [], []

    asserted = [(c, el) for c, el, st in rec.sequence if st == "asserted"]
    before = [(c, el) for c, el in asserted if tc is None or c <= tc + 0.01]
    prot = [(c, el, element_class(el)) for c, el in before if element_class(el) not in NON_PROTECTION]
    initiators = [x for x in prot if tc is not None and tc - x[0] <= 0.5]
    pickups = [x for x in prot if x not in initiators]

    eq_hits = {}
    for key, _ in TRIP_EQUATIONS:
        if settings.get(key):
            hits = equation_hits(settings[key], [el for _, el in before])
            if hits:
                eq_hits[key] = hits

    # ---- headline and evidence
    if tripped:
        cause.append(f"Relay TRIPPED. {desc or 'Event'}.")
    else:
        cause.append(f"No trip recorded. {desc}." if desc else "No trip recorded for this event.")
    if tdesc:
        cause.append("Front-panel targets indicate: " + ", ".join(tdesc) + f"  [{rec.targets.strip()}]")

    primary = []
    if "TR" in eq_hits:
        primary = eq_hits["TR"]
        cause.append("Trip equation (TR) elements that asserted before the trip: " + ", ".join(primary)
                     + " -> these caused the trip.")
    comm_bits = [el for _, el in before if element_class(el) == "communications / transfer trip"]
    for key, what in TRIP_EQUATIONS[1:]:
        if key in eq_hits:
            primary = primary or eq_hits[key]
            line = f"Elements in the {what} ({key}) asserted before the trip: " + ", ".join(eq_hits[key])
            if key == "TRCOMM" and comm_bits:
                line += (f", with {', '.join(comm_bits)} from the communications scheme -> "
                         "communications-assisted trip (e.g. POTT/DCUB)")
                primary = primary + [b for b in comm_bits if b not in primary and b != "KEY"]
            cause.append(line + ".")
    if not primary and initiators:
        primary = [el for _, el, _ in initiators]
        cause.append("Element(s) that asserted with the trip: "
                     + ", ".join(f"{el} ({cls}, {c:.2f} cyc)" for c, el, cls in initiators)
                     + " -> most likely tripping element(s).")
    if pickups:
        cause.append("Earlier pickups before the trip: "
                     + ", ".join(f"{el} ({cls}, {c:.2f} cyc)" for c, el, cls in pickups[:10]))

    classes = {cls for _, _, cls in initiators} | {element_class(el) for el in primary}
    if tripped and not is_fault:
        if classes & {"communications / transfer trip", "contact input"}:
            cause.append("No fault type was identified and a communications/input bit asserted with the trip: "
                         "most likely a direct/remote transfer trip or an external trip input, not a local fault.")
        elif "frequency" in classes:
            cause.append("Frequency element tripped: likely underfrequency load shedding or a system disturbance.")
        elif classes & {"undervoltage", "overvoltage"}:
            cause.append("Voltage element tripped: check system voltage conditions or loss of source.")
        elif "breaker failure" in classes:
            cause.append("Breaker-failure logic operated: a breaker failed to clear a fault.")
        elif "out-of-step" in classes:
            cause.append("Out-of-step logic operated: power swing / loss of synchronism.")
        elif "COMM" in (rec.targets or "").upper().split():
            cause.append("COMM target with no fault type: most likely a transfer trip received from the remote end.")
        elif re.search(r"\b81\b", rec.targets or ""):
            cause.append("Frequency (81) target: likely underfrequency load shedding or a system disturbance.")

    if tc is not None and rec.open_cycle is not None and rec.open_cycle > tc:
        cause.append(f"Breaker opened ~{rec.open_cycle - tc:.2f} cycles after trip "
                     f"(trip at {tc:.2f} cyc, 52A dropped at {rec.open_cycle:.2f} cyc).")
    if rec.shot and rec.shot not in ("0", ""):
        cause.append(f"Reclose shot counter = {rec.shot} (event occurred during a reclose sequence).")
    cause.extend(rec.notes)

    # ---- what we can't know from this data (gaps)
    if rec.kind == "history" and tripped and not rec.dup_of:
        gaps.append("History row only, so the cause is read from the targets. "
                    "Load this event's .CEV report to confirm which element tripped.")
    elif rec.kind != "history" and not rec.sequence and tripped:
        gaps.append("No relay element data in this file, so the tripping element is inferred from the "
                    "front-panel targets. Load the .CEV report for this event to confirm.")
    if tripped and not settings.get("TR"):
        if how == "mismatch":
            gaps.append(f"Settings files are loaded, but none match relay '{rec.relay_id}', so no settings "
                        "were applied. Add this relay's .RDB.")
        else:
            gaps.append("No settings for this relay, so the trip equation (TR) and line length couldn't be "
                        "checked. Add the relay's .RDB (or use a CEV that includes settings).")
    if settings.get("__method__") == "scan":
        gaps.append("Settings were read with a best-effort scan of the .RDB file; some values may be missing.")

    # ---- things that look wrong or ambiguous (checks)
    if how == "assumed":
        checks.append(f"This event has no relay ID, so settings from the only loaded RDB ({label}) were "
                      "assumed to apply. Confirm it's the same relay.")
    if rec.event_type and core not in FAULT_TYPES and core not in NON_FAULT_TYPES:
        checks.append(f"Event type '{rec.event_type}' isn't a code this tool recognizes; interpret it manually.")
    if core == "TRIP":
        checks.append("The relay tripped without identifying a fault type. Confirm whether this was a transfer "
                      "trip, a manual/SCADA open, or an external trip input.")

    ep, tp = event_phases(rec.event_type), target_phases(rec.targets)
    if ep and tp:
        e_abc, t_abc = ep - {"G"}, tp - {"G"}
        if t_abc and e_abc != t_abc:
            checks.append(f"Targets show phase(s) {''.join(sorted(t_abc))} but the event type says "
                          f"{''.join(sorted(e_abc))}.")
        if t_abc and ("G" in ep) != ("G" in tp):
            checks.append("Targets and event type disagree on whether ground was involved.")

    ph, gnd = phase_currents(rec)
    if ep and len(ph) >= 2:
        top = max(ph, key=lambda k: ph[k][1])
        top_pre, top_fault = ph[top]
        e_abc = ep - {"G"}
        if top not in e_abc:
            checks.append(f"Highest fault current is on phase {top} ({top_fault:,.0f} A), but the event type "
                          f"says {rec.event_type.split()[0]}.")
        if top_pre and top_fault < 1.5 * top_pre:
            checks.append(f"Fault current barely rose above load ({top_fault:,.0f} A vs {top_pre:,.0f} A "
                          "prefault). Could be a high-impedance fault, a weak source, or a misclassified event.")
        if gnd is not None and top_fault > 0:
            g = gnd[1]
            if "G" in ep and g < 0.1 * top_fault:
                checks.append(f"Ground fault reported, but ground current ({g:,.0f} A) is small next to phase "
                              f"current ({top_fault:,.0f} A).")
            if "G" not in ep and g > 0.3 * top_fault:
                checks.append(f"Significant ground current ({g:,.0f} A) for a fault classified without ground "
                              f"({rec.event_type.split()[0]}).")

    loc = (rec.location or "").strip()
    v = fnum(loc)
    ll = fnum(settings.get("LL", ""))
    if is_fault:
        if not loc:
            gaps.append("No fault location in this file.")
        elif "$" in loc:
            checks.append("The relay couldn't calculate a fault location ($$$$$). This usually means too "
                          "little fault data or a very fast clearing.")
    if v is not None and is_fault:
        if v < 0:
            checks.append(f"Negative fault location ({loc}) usually means the fault was behind the relay "
                          "(reverse direction). Check the directional elements and CT polarity.")
        elif ll:
            pct = 100 * v / ll
            if v > ll * 1.02:
                checks.append(f"Fault location {loc} is beyond the end of the line (LL = {settings['LL']}). "
                              "Possible fault on the adjacent section, or a location error.")
            zone1 = "Zone 1 distance" in tdesc or any(re.match(r"^(Z|M)1", el) for el in primary)
            if zone1 and pct > 90:
                checks.append(f"Zone 1 operated for a fault at ~{pct:.0f}% of the line, beyond typical Zone 1 "
                              "reach (80-90%). Check for overreach or a location error.")
        elif settings.get("TR") or not tripped:  # (missing settings already flagged above)
            gaps.append("Line length (LL) isn't known, so the fault location can't be checked against the "
                        "line. Load the relay's .RDB.")

    f = fnum(rec.freq)
    if f:
        nominal = 50.0 if abs(f - 50) < abs(f - 60) else 60.0
        if abs(f - nominal) > 0.5:
            checks.append(f"Frequency was {f:.2f} Hz, well off nominal {nominal:.0f} Hz.")

    if rec.sequence:
        if tripped_by_type and tc is None:
            checks.append("The event type shows a trip (T), but no TRIP bit asserted in this record.")
        if tc is not None and rec.event_type and not tripped_by_type:
            checks.append("TRIP asserted in the record, but the event type doesn't show a trip.")
        if tripped and settings.get("TR") and not eq_hits:
            checks.append("No element from the trip equations asserted before the trip. It may have come "
                          "from a manual/SCADA command, another logic path, or an element not in this record.")
        if tripped and not prot and not eq_hits:
            checks.append("No protection element asserted before the trip in this record.")
        if tc is not None:
            if rec.open_cycle is None:
                checks.append("No breaker-open indication in the record. Confirm the breaker opened (the record "
                              "may end too soon, or 52A isn't in this report).")
            elif rec.open_cycle < tc:
                checks.append("Breaker status dropped before the trip asserted. The breaker may have been "
                              "opened by another relay or manually.")
            elif rec.open_cycle - tc > 5:
                checks.append(f"Breaker took {rec.open_cycle - tc:.1f} cycles to open after the trip "
                              "(typical is 2-5). Check the breaker and trip circuit.")
        if rec.prefault_short:
            checks.append("The record triggered very early, so 'prefault' values may include fault data.")

    # ---- confidence
    if rec.kind == "history":
        level = 1
    elif rec.sequence and settings.get("TR"):
        level = 3
    elif rec.sequence or rec.targets:
        level = 2
    else:
        level = 1
    if checks:
        level -= 2 if len(checks) >= 3 else 1
    level = max(1, level)
    basis = ("history row" if rec.kind == "history" else
             "element data" if rec.sequence else
             "targets only" if rec.targets else "event type only")
    if settings.get("TR"):
        basis += ", settings" + (f" from {label}" if label else " in event file")
    return {"tripped": tripped, "desc": desc, "cause": cause, "checks": checks, "gaps": gaps,
            "confidence": {3: "High", 2: "Medium", 1: "Low"}[level], "basis": basis, "primary": primary}


def determine_cause(rec, settings, how="none", label=""):
    a = analyze_event(rec, settings, how, label)
    return a["cause"], a["tripped"]


def short_cause(rec, settings, analysis=None):
    a = analysis or analyze_event(rec, settings)
    parts = []
    if a["desc"]:
        d = re.sub(r"\s*\((?:non-fault|phase-to-phase|event report)[^)]*\)", "", a["desc"]).split(" (")[0]
        parts.append(re.sub(r" - no trip$", "", d))
    if a["primary"]:
        parts.append("via " + ", ".join(a["primary"][:3]))
    else:
        td = interpret_targets(rec.targets)
        if td:
            parts.append("via " + "/".join(td[:3]))
    loc = rec.location if rec.location and "$" not in rec.location else ""
    if loc:
        parts.append(f"@ {loc}")
    n = len(a["checks"]) + len(a["gaps"])
    tail = f"  [{n} to check]" if n else ""
    return ("TRIP - " if a["tripped"] else "No trip - ") + " ".join(parts) + tail


def event_datetime(rec):
    """Best-effort datetime for sorting and duplicate matching."""
    d = (rec.date or "").strip()
    t = (rec.time or "").strip()
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            day = datetime.datetime.strptime(d, fmt)
            break
        except ValueError:
            continue
    else:
        return None
    m = re.match(r"(\d{1,2}):(\d{2}):(\d{2})(?:\.(\d+))?", t)
    if m:
        frac = float("0." + m.group(4)) if m.group(4) else 0.0
        day = day.replace(hour=int(m.group(1)), minute=int(m.group(2)), second=int(m.group(3)),
                          microsecond=int(frac * 1e6))
    return day


def link_duplicates(events):
    """Mark history rows that are the same event as a detailed report (dup_of = that report)."""
    detailed = [(e, event_datetime(e)) for e in events if e.kind == "event"]
    for e in events:
        if e.kind != "history":
            continue
        dt = event_datetime(e)
        if not dt:
            continue
        for d, ddt in detailed:
            if ddt and abs((ddt - dt).total_seconds()) <= 1.0 and \
                    (not e.relay_id or not d.relay_id or e.relay_id.upper() == d.relay_id.upper()):
                e.dup_of = d
                break


# --------------------------------------------------------------------------
# View options: filter, sort and detail level
# --------------------------------------------------------------------------
SHOW_OPTIONS = (("all", "All events"), ("trips", "Trips only"), ("review", "Needs review"),
                ("faults", "Faults only"), ("notrip", "No-trip events"))
SORT_OPTIONS = (("time", "Oldest first"), ("newest", "Newest first"), ("relay", "By relay"),
                ("review", "Needs review first"))
DETAIL_OPTIONS = (("summary", "Summary only"), ("standard", "Standard"), ("full", "Full"))
FAULT_GROUPS = ((("single phase-to-ground",) * 2, ("AG", "BG", "CG")),
                (("phase-to-phase",) * 2, ("AB", "BC", "CA")),
                (("two-phase-to-ground",) * 2, ("ABG", "BCG", "CAG")),
                (("three-phase",) * 2, ("ABC", "ABCG")),
                (("trip without a fault type", "trips without a fault type"), ("TRIP",)))


def relay_key(e):
    return e.relay_id or e.terminal_id or e.source


def show_matches(e, a, show):
    if show == "trips":
        return a["tripped"]
    if show == "review":
        return bool(a["checks"] or a["gaps"])
    if show == "faults":
        return event_core(e.event_type) in FAULT_TYPES
    if show == "notrip":
        return not a["tripped"]
    return True


def search_text(e, a):
    parts = [e.relay_id, e.terminal_id, e.source, e.event_type, e.targets, e.date, e.time, e.location,
             a["desc"]] + a["cause"] + a["checks"] + a["gaps"] + a["primary"]
    return " ".join(p for p in parts if p).upper()


def prepare_view(events, rdb, opts=None):
    """Analyze every event, then apply the view options.

    Returns (items, shown, n_dups_hidden). Each item is a dict with e, s (settings),
    label, how, a (analysis) and order (position in time order)."""
    opts = opts or {}
    ev = [e for e in events if e.kind in ("event", "history")]
    ev.sort(key=lambda e: (event_datetime(e) is None, event_datetime(e) or datetime.datetime.min))
    for e in ev:
        e.notes = [n for n in e.notes if not n.startswith("Same event as")]
        e.dup_of = None
    link_duplicates(ev)
    items = []
    for i, e in enumerate(ev):
        s_rdb, label, how = match_settings(e, rdb)
        s = {**e.settings, **s_rdb}
        items.append({"e": e, "s": s, "label": label, "how": how, "a": analyze_event(e, s, how, label), "order": i})

    show = opts.get("show") or "all"
    relay = (opts.get("relay") or "").strip().upper()
    search = (opts.get("search") or "").strip().upper()
    hide_dups = bool(opts.get("hide_duplicates"))
    shown, n_dups = [], 0
    for it in items:
        e, a = it["e"], it["a"]
        if not show_matches(e, a, show):
            continue
        if relay and relay_key(e).strip().upper() != relay:
            continue
        if search and search not in search_text(e, a):
            continue
        if hide_dups and e.dup_of is not None:
            n_dups += 1
            continue
        shown.append(it)

    sort = opts.get("sort") or "time"
    if sort == "newest":
        dated = [it for it in shown if event_datetime(it["e"]) is not None]
        shown = dated[::-1] + [it for it in shown if event_datetime(it["e"]) is None]
    elif sort == "relay":
        shown.sort(key=lambda it: (relay_key(it["e"]).upper(), it["order"]))
    elif sort == "review":
        shown.sort(key=lambda it: (-(len(it["a"]["checks"]) + len(it["a"]["gaps"])), it["order"]))
    return items, shown, n_dups


def view_choices(events, rdb):
    """Counts for the Show options and the relays present, for labelling the view controls."""
    items, _, _ = prepare_view(events, rdb)
    counts = {key: sum(1 for it in items if show_matches(it["e"], it["a"], key)) for key, _ in SHOW_OPTIONS}
    relays = {}
    for it in items:
        k = relay_key(it["e"])
        relays[k] = relays.get(k, 0) + 1
    return {"counts": counts, "relays": sorted(relays.items(), key=lambda kv: kv[0].upper())}


def describe_view(opts, n_dups):
    opts = opts or {}
    parts = []
    show = opts.get("show") or "all"
    if show != "all":
        parts.append(dict(SHOW_OPTIONS)[show].lower())
    if (opts.get("relay") or "").strip():
        parts.append(f"relay {opts['relay'].strip()}")
    sort = opts.get("sort") or "time"
    if sort != "time":
        parts.append(dict(SORT_OPTIONS)[sort].lower())
    if (opts.get("search") or "").strip():
        parts.append(f'matching "{opts["search"].strip()}"')
    if n_dups:
        parts.append(f"{n_dups} duplicate history row{'s' if n_dups != 1 else ''} hidden")
    return "; ".join(parts)


def overview_lines(shown):
    a_list = [it["a"] for it in shown]
    trips = sum(1 for a in a_list if a["tripped"])
    lines = [f"  Events      : {len(shown)} - {trips} trip{'s' if trips != 1 else ''}, {len(shown) - trips} no trip"]
    cats, known = [], set()
    for (one, many), codes in FAULT_GROUPS:
        n = sum(1 for it in shown if event_core(it["e"].event_type) in codes)
        known.update(codes)
        if n:
            cats.append(f"{n} {one if n == 1 else many}")
    other = sum(1 for it in shown if event_core(it["e"].event_type) not in known)
    if other:
        cats.append(f"{other} other event{'s' if other != 1 else ''} (no fault)")
    lines.append("  Fault types : " + ", ".join(cats))
    relays = {}
    for it in shown:
        k = relay_key(it["e"])
        relays[k] = relays.get(k, 0) + 1
    ordered = sorted(relays.items(), key=lambda kv: (-kv[1], kv[0].upper()))
    lines.append("  Relays      : " + ", ".join(f"{k} ({n})" for k, n in ordered))
    conf = {c: sum(1 for a in a_list if a["confidence"] == c) for c in ("High", "Medium", "Low")}
    lines.append("  Confidence  : " + ", ".join(f"{n} {c}" for c, n in conf.items() if n))
    return lines


def build_report(events, rdb, errors, files, comments="", opts=None):
    opts = opts or {}
    detail = opts.get("detail") or "full"
    show_seq = detail == "full" and opts.get("sequence", True)
    show_settings = detail == "full" and opts.get("settings", True)
    show_ser = detail == "full" and opts.get("ser", True)
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    items, shown, n_dups = prepare_view(events, rdb, opts)
    ser = [e for e in events if e.kind == "ser"]
    n_items = sum(len(it["a"]["checks"]) + len(it["a"]["gaps"]) for it in shown)
    n_ev = sum(1 for it in shown if it["a"]["checks"] or it["a"]["gaps"])
    number = {id(it["e"]): n for n, it in enumerate(shown, 1)}

    def dup_note(e):
        d = e.dup_of
        if d is None:
            return []
        if id(d) in number:
            return [f"Same event as #{number[id(d)]} ({d.source}), which has more detail."]
        return [f"Same event as the detailed report in {d.source} (not shown in this view)."]

    out = []
    out.append("SEL TRIP EVENT SUMMARY")
    out.append("=" * 60)
    out.append(f"Generated: {now}")
    out.append(f"Files analyzed: {len(files)}   Events found: {len(items)}   "
               f"Settings databases: {len(rdb)}")
    view = describe_view(opts, n_dups)
    if items and (view or len(shown) != len(items)):
        out.append(f"Showing: {len(shown)} of {len(items)} events" + (f" ({view})" if view else ""))
    if shown:
        out.append(f"Needs review: {n_items} item(s) to check across {n_ev} of {len(shown)} event(s)"
                   + (" - see the CHECK lines" if n_items else ""))
    out.append("")

    if comments.strip():
        out.append("ANALYST COMMENTS / NOTES")
        out.append("-" * 60)
        out.extend(comments.strip().splitlines())
        out.append("")

    if not items and not ser:
        out.append("No event data recognized in the loaded files.")
    else:
        if len(shown) >= 2:
            out.append("OVERVIEW")
            out.append("-" * 60)
            out.extend(overview_lines(shown))
            out.append("")

        out.append("QUICK SUMMARY")
        out.append("-" * 60)
        if items and not shown:
            out.append("No events match the current Show / Relay / Search options.")
        for n, it in enumerate(shown, 1):
            e, s, a = it["e"], it["s"], it["a"]
            when = f"{e.date} {e.time}".strip() or "date n/a"
            out.append(f"{n:>2}. {when} | {relay_key(e)} | {e.event_type or '?'} | {short_cause(e, s, a)}")
        out.append("")

        if detail != "summary" and not (items and not shown):
            out.append("EVENT DETAILS")
            out.append("-" * 60)
            shown_settings = {}
            for n, it in enumerate(shown, 1):
                e, s, label, a = it["e"], it["s"], it["label"], it["a"]
                out.append(f"Event {n} of {len(shown)}  ({'history row' if e.kind == 'history' else 'event report'})")
                out.append(f"  Source file : {e.source}")
                relay = " / ".join(x for x in (e.relay_id, e.terminal_id) if x)
                if relay:
                    out.append(f"  Relay       : {relay}")
                if e.fid:
                    out.append(f"  Firmware    : {e.fid}")
                if e.date or e.time:
                    out.append(f"  Date / Time : {e.date} {e.time}")
                if e.event_type:
                    out.append(f"  Event type  : {e.event_type}  ({a['desc']})")
                if e.location:
                    out.append(f"  Fault loc.  : {fmt_loc(e.location, s)}")
                if e.targets:
                    out.append(f"  Targets     : {e.targets}")
                if e.freq:
                    out.append(f"  Frequency   : {e.freq} Hz")
                if e.shot:
                    out.append(f"  Shot        : {e.shot}")
                for x in e.extra:
                    out.append(f"  {x}")
                if e.currents:
                    parts = []
                    for k, (pre, post) in e.currents.items():
                        parts.append(f"{k} {post:,.0f}" + (f" (pre {pre:,.0f})" if pre is not None else ""))
                    est = ", estimated from samples" if e.sequence or any(p is not None for p, _ in e.currents.values()) else ""
                    out.append("  Currents    : " + "; ".join(parts) + f"  [A, max during fault{est}]")
                if e.voltages:
                    parts = [f"{k} {pre:,.2f} -> {post:,.2f}" for k, (pre, post) in e.voltages.items()]
                    out.append("  Voltages    : " + "; ".join(parts) + "  [prefault -> min during fault]")
                n_chk = len(a["checks"]) + len(a["gaps"])
                out.append(f"  Confidence  : {a['confidence']} ({a['basis']})"
                           + (f" - {n_chk} item(s) to check" if n_chk else ""))

                out.append("  Probable cause:")
                for c in a["cause"] + dup_note(e):
                    out.append(f"    - {c}")
                if n_chk:
                    out.append("  Check before relying on this:")
                    for c in a["checks"] + a["gaps"]:
                        out.append(f"    - CHECK: {c}")

                if show_seq and e.sequence:
                    out.append("  Element sequence (cycles from start of record):")
                    rel = [x for x in e.sequence if element_class(x[1])]
                    rest = len(e.sequence) - len(rel)
                    for cyc, el, st in rel[:40]:
                        tag = "already asserted at start" if st == "pre" else "asserted"
                        out.append(f"    {cyc:6.2f}  {el:<10} {tag}  ({element_class(el)})")
                    if len(rel) > 40:
                        out.append(f"    ... {len(rel) - 40} more")
                    if rest:
                        out.append(f"    ({rest} other relay word bits changed - not listed)")

                if show_settings and s:
                    keys = [k for k in ("TR", "TRCOMM", "TRSOTF", "DTT", "LL", "CTR", "PTR", "Z1MAG", "Z1P",
                                        "Z1MG", "50P1P", "50G1P", "51PP", "51GP", "ER") if k in s]
                    if keys:
                        src = f"RDB: {label}" if label else "event file"
                        sig = (src, tuple((k, s[k]) for k in keys))
                        if sig in shown_settings:
                            out.append(f"  Settings    : same as Event {shown_settings[sig]} ({src})")
                        else:
                            shown_settings[sig] = n
                            out.append(f"  Settings context ({src}):")
                            for k in keys:
                                out.append(f"    {k} := {s[k]}")
                out.append("")

    if show_ser and ser:
        out.append("SEQUENTIAL EVENTS RECORDER (trip-related lines)")
        out.append("-" * 60)
        for r in ser:
            out.append(f"  From {r.source}:")
            lines = [ln for ln in r.ser_lines if element_class(ln.split()[2])] or r.ser_lines
            for ln in lines[:80]:
                out.append(f"    {ln}")
        out.append("")

    if rdb and not items:
        out.append("SETTINGS LOADED")
        out.append("-" * 60)
        for label, s in rdb.items():
            out.append(f"  {label}: TR := {s.get('TR', 'n/a')}")
        out.append("")

    if errors:
        out.append("FILES THAT COULD NOT BE READ")
        out.append("-" * 60)
        out.extend("  " + e for e in errors)
        out.append("")

    out.append("Note: automated interpretation - verify against the relay oscillography and settings.")
    return "\n".join(out)


CSV_COLUMNS = ["#", "Date", "Time", "Relay", "Station", "Event type", "Description", "Tripped",
               "Tripping elements", "Fault location", "% of line", "Targets", "Max phase current (A)",
               "Ground current (A)", "Breaker clearing (cycles)", "Confidence", "Items to check",
               "Check details", "Source file"]


def report_csv(events, rdb, opts=None):
    """One row per event (same filter and order as the summary) for Excel."""
    import io
    _, shown, _ = prepare_view(events, rdb, opts)
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(CSV_COLUMNS)
    for n, it in enumerate(shown, 1):
        e, s, a = it["e"], it["s"], it["a"]
        ph, gnd = phase_currents(e)
        peaks = [v[1] for v in ph.values()] or [v[1] for k, v in e.currents.items() if k == "Fault current"]
        loc = (e.location or "").strip()
        v, ll = fnum(loc), fnum(s.get("LL", ""))
        clearing = (f"{e.open_cycle - e.trip_cycle:.2f}" if e.trip_cycle is not None and e.open_cycle is not None
                    and e.open_cycle > e.trip_cycle else "")
        w.writerow([str(n), e.date, e.time, e.relay_id, e.terminal_id, e.event_type, a["desc"],
                    "Yes" if a["tripped"] else "No", ", ".join(a["primary"]), loc,
                    f"{100 * v / ll:.0f}" if v is not None and ll else "", e.targets,
                    f"{max(peaks):.0f}" if peaks else "", f"{gnd[1]:.0f}" if gnd is not None else "",
                    clearing, a["confidence"], str(len(a["checks"]) + len(a["gaps"])),
                    " | ".join(a["checks"] + a["gaps"]), e.source])
    return buf.getvalue()


# --------------------------------------------------------------------------
# Report export (Word / PDF) - built from the summary text *as edited*
# --------------------------------------------------------------------------
REPORT_COLORS = {"ink": "1B1B1B", "muted": "666666", "accent": "0B5CAD", "rule": "D6DCE4",
                 "head_bg": "EAF1FA", "trip": "C42B1C", "ok": "0F7B0F", "band": "F7F9FC",
                 "warn": "A15C00"}

KV_RE = re.compile(r"^ {2}([A-Za-z][A-Za-z ./]{1,14}?)\s*:\s?(.*)$")
SUB_RE = re.compile(r"^ {2}(\S.*?):\s*$")
SEQ_RE = re.compile(r"^\s{3,}([\d.]+)\s+(\S+)\s+(already asserted at start|asserted)\s+\((.*)\)\s*$")
QS_RE = re.compile(r"^\s*(\d+)\.\s+(.*)$")


def parse_report_text(text):
    """Turn the plain-text summary into typed blocks for the Word/PDF writers.
    Anything unrecognised (e.g. lines the user typed) becomes a paragraph."""
    blocks, para = [], []
    section, sub = "", ""
    lines = text.replace("\r\n", "\n").split("\n")

    def flush():
        if para:
            blocks.append(("para", "\n".join(para)))
            para.clear()

    def add(kind, *payload):
        flush()
        if blocks and blocks[-1][0] == kind and kind in ("kv", "bullets", "mono", "meta"):
            blocks[-1][1].extend(payload[0])
        elif blocks and kind == "table" and blocks[-1][0] == "table" and blocks[-1][1] == payload[0]:
            blocks[-1][2].extend(payload[1])
        else:
            blocks.append((kind,) + tuple(list(p) if isinstance(p, list) else p for p in payload))

    for idx, ln in enumerate(lines):
        s = ln.strip()
        if not s:
            flush()
            continue
        if re.fullmatch(r"[=\-]{6,}", s):
            continue
        if not blocks and not para and s.upper().startswith("SEL TRIP"):
            blocks.append(("title", s.title().replace("Sel ", "SEL ")))
            continue
        if not section and re.match(r"^(Generated|Files analyzed|Showing|Needs review)\b", s):
            add("meta", [s])
            continue
        if ln == s and re.fullmatch(r"[A-Z][A-Z /&\-]{3,}(\s*\(.*\))?", s):
            section, sub = s, ""
            add("h1", s.title().replace("Ser", "SER").replace("Rdb", "RDB").replace("/ ", "/ "))
            continue
        if s.startswith("Note:") and ln == s:
            add("note", s)
            continue
        if section == "QUICK SUMMARY":
            m = QS_RE.match(ln)
            if m and m.group(2).count("|") >= 3:
                parts = [p.strip() for p in m.group(2).split("|", 3)]
                add("table", ["#", "Date / time", "Relay", "Type", "Cause"], [[m.group(1)] + parts])
                continue
        if re.match(r"^Event \d+ of \d+", ln):
            sub = ""
            add("h2", s)
            continue
        m = SUB_RE.match(ln)
        if m and not KV_RE.match(ln) or (m and m.group(1).startswith(("Probable", "Element", "Settings", "From"))):
            sub = m.group(1)
            add("sub", sub)
            continue
        if ln.startswith("    "):  # content under a sub-heading
            if s.startswith("- "):
                add("bullets", [s[2:]])
                continue
            m = SEQ_RE.match(ln)
            if m:
                add("table", ["Cycle", "Element", "State", "Function"],
                    [[m.group(1), m.group(2), m.group(3), m.group(4)]])
                continue
            m = re.match(r"^\s+([A-Z0-9_]+) := (.*)$", ln)
            if m:
                add("kv", [(m.group(1), m.group(2), True)])
                continue
            if sub.startswith("From") or section.startswith("SEQUENTIAL"):
                add("mono", [s])
                continue
            add("note", s)
            continue
        m = KV_RE.match(ln)
        if m and section in ("EVENT DETAILS", "OVERVIEW"):
            add("kv", [(m.group(1).strip(), m.group(2).strip(), False)])
            continue
        if section.startswith("FILES THAT") and ln.startswith("  "):
            add("bullets", [s])
            continue
        para.append(s)
    flush()
    return blocks


def _trip_runs(text):
    """Split a cause string so the TRIP / No trip prefix can be coloured."""
    runs = [(text, None)]
    for word, color in (("Relay TRIPPED", "trip"), ("TRIP", "trip"), ("No trip recorded", "ok"),
                        ("No trip", "ok"), ("CHECK:", "warn")):
        if text.startswith(word):
            runs = [(word, color), (text[len(word):], None)]
            break
    m = re.search(r"\s*\[\d+ to check\]$", runs[-1][0])
    if m:  # quick-summary suffix
        last = runs[-1][0]
        runs[-1] = (last[:m.start()], runs[-1][1])
        runs.append((last[m.start():], "warn"))
    return runs


def export_docx(text, path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    C = {k: RGBColor.from_string(v) for k, v in REPORT_COLORS.items()}
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    for name, size, color in (("Heading 1", 14, C["accent"]), ("Heading 2", 12, C["ink"]),
                              ("Heading 3", 10.5, C["muted"])):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.color.rgb, st.font.bold = "Calibri", Pt(size), color, True
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(4)
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(a), "Calibri")

    def shade(cell, hexcolor):
        tcpr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), hexcolor)
        tcpr.append(sh)

    def borders(table, color=REPORT_COLORS["rule"], inside=True):
        tblpr = table._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            on = edge in ("top", "bottom", "insideH") and (inside or edge != "insideH")
            e.set(qn("w:val"), "single" if on else "nil")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), color)
            b.append(e)
        tblpr.append(b)

    def set_widths(t, inches):
        t.autofit = False
        tblpr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblpr.append(lay)
        for col, w in zip(t.columns, inches):
            col.width = Inches(w)
        for row in t.rows:
            for c, w in zip(row.cells, inches):
                c.width = Inches(w)

    def keep_together(t):
        if len(t.rows) > 18:
            return
        for row in t.rows[:-1]:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True

    def cell_text(cell, txt, bold=False, color=None, mono=False, size=9.5, runs=None):
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        for t, col in (runs or [(txt, None)]):
            r = p.add_run(t)
            r.bold = bold or col is not None
            r.font.size = Pt(size)
            if col or color:
                r.font.color.rgb = C[col] if col else color
            if mono:
                r.font.name = "Consolas"

    def field(par, instr):
        r = par.add_run()
        for tag, txt in (("begin", None), (None, instr), ("end", None)):
            if tag:
                f = OxmlElement("w:fldChar")
                f.set(qn("w:fldCharType"), tag)
                r._r.append(f)
            else:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = txt
                r._r.append(it)
        return r

    # footer: app + page x of y
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for t in (f"{APP_NAME} v{VERSION}    Page ",):
        r = fp.add_run(t)
        r.font.size, r.font.color.rgb = Pt(8), C["muted"]
    for instr, tail in (("PAGE", " of "), ("NUMPAGES", "")):
        field(fp, instr).font.size = Pt(8)
        if tail:
            r = fp.add_run(tail)
            r.font.size, r.font.color.rgb = Pt(8), C["muted"]

    for blk in parse_report_text(text):
        kind = blk[0]
        if kind == "title":
            p = doc.add_paragraph()
            r = p.add_run(blk[1])
            r.bold, r.font.size, r.font.color.rgb = True, Pt(22), C["ink"]
            p.paragraph_format.space_after = Pt(2)
        elif kind == "meta":
            p = doc.add_paragraph()
            r = p.add_run("    ".join(blk[1]))
            r.font.size, r.font.color.rgb = Pt(9), C["muted"]
            p.paragraph_format.space_after = Pt(10)
        elif kind == "h1":
            doc.add_heading(blk[1], level=1)
        elif kind == "h2":
            doc.add_heading(blk[1], level=2)
        elif kind == "sub":
            doc.add_heading(blk[1], level=3)
        elif kind == "para":
            p = doc.add_paragraph()
            for k, ln in enumerate(blk[1].split("\n")):
                if k:
                    p.add_run().add_break()
                p.add_run(ln)
        elif kind == "note":
            p = doc.add_paragraph()
            r = p.add_run(blk[1])
            r.italic, r.font.size, r.font.color.rgb = True, Pt(9), C["muted"]
        elif kind == "bullets":
            for b in blk[1]:
                p = doc.add_paragraph(style="List Bullet")
                for t, col in _trip_runs(b):
                    r = p.add_run(t)
                    if col:
                        r.bold, r.font.color.rgb = True, C[col]
        elif kind == "mono":
            p = doc.add_paragraph()
            for k, ln in enumerate(blk[1]):
                if k:
                    p.add_run().add_break()
                r = p.add_run(ln)
                r.font.name, r.font.size = "Consolas", Pt(8.5)
        elif kind == "kv":
            t = doc.add_table(rows=0, cols=2)
            borders(t)
            for k, v, mono in blk[1]:
                row = t.add_row().cells
                cell_text(row[0], k, bold=True, color=C["muted"], mono=mono)
                cell_text(row[1], v, mono=mono)
            set_widths(t, [1.5, 5.4])
            keep_together(t)
        elif kind == "table":
            headers, rows = blk[1], blk[2]
            t = doc.add_table(rows=1, cols=len(headers))
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            borders(t)
            for c, h in zip(t.rows[0].cells, headers):
                shade(c, REPORT_COLORS["head_bg"])
                cell_text(c, h, bold=True, color=C["accent"], size=9)
            for i, r in enumerate(rows):
                cells = t.add_row().cells
                for j, (c, v) in enumerate(zip(cells, r)):
                    if i % 2:
                        shade(c, REPORT_COLORS["band"])
                    is_cause = headers[j] == "Cause"
                    cell_text(c, v, size=9, runs=_trip_runs(v) if is_cause else None,
                              mono=headers[0] == "Cycle" and j < 2)
            if headers[0] == "#":
                widths = [0.35, 1.45, 1.0, 0.6, 3.5]
            elif headers[0] == "Cycle":
                widths = [0.7, 1.1, 1.8, 3.3]
            else:
                widths = [6.9 / len(headers)] * len(headers)
            set_widths(t, widths)
            keep_together(t)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
    doc.save(path)


def export_pdf(text, path):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (KeepTogether, ListFlowable, ListItem, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)

    H = {k: colors.HexColor("#" + v) for k, v in REPORT_COLORS.items()}

    def ps(name, **kw):
        base = dict(fontName="Helvetica", fontSize=9.5, leading=12.5, textColor=H["ink"], alignment=TA_LEFT)
        base.update(kw)
        return ParagraphStyle(name, **base)

    S = {
        "title": ps("t", fontName="Helvetica-Bold", fontSize=20, leading=24, spaceAfter=2),
        "meta": ps("m", fontSize=8.5, textColor=H["muted"], spaceAfter=10),
        "h1": ps("h1", keepWithNext=1, fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=H["accent"],
                 spaceBefore=12, spaceAfter=5),
        "h2": ps("h2", keepWithNext=1, fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=10, spaceAfter=4),
        "sub": ps("h3", keepWithNext=1, fontName="Helvetica-Bold", fontSize=9.5, textColor=H["muted"], spaceBefore=7, spaceAfter=3),
        "body": ps("b", spaceAfter=4),
        "note": ps("n", fontName="Helvetica-Oblique", fontSize=8.5, textColor=H["muted"], spaceAfter=4),
        "cell": ps("c", fontSize=8.5, leading=10.5),
        "cellb": ps("cb", fontName="Helvetica-Bold", fontSize=8.5, leading=10.5, textColor=H["muted"]),
        "head": ps("hd", fontName="Helvetica-Bold", fontSize=8.5, leading=10.5, textColor=H["accent"]),
        "mono": ps("mo", fontName="Courier", fontSize=8, leading=10),
    }

    def rich(t):
        out = ""
        for part, col in _trip_runs(t):
            e = escape(part)
            out += f'<font color="#{REPORT_COLORS[col]}"><b>{e}</b></font>' if col else e
        return out

    width = letter[0] - 1.6 * inch
    story = []
    for blk in parse_report_text(text):
        kind = blk[0]
        if kind == "title":
            story.append(Paragraph(escape(blk[1]), S["title"]))
        elif kind == "meta":
            story.append(Paragraph(escape("    ".join(blk[1])), S["meta"]))
        elif kind in ("h1", "h2", "sub"):
            story.append(Paragraph(escape(blk[1]), S[kind]))
        elif kind == "para":
            story.append(Paragraph(escape(blk[1]).replace("\n", "<br/>"), S["body"]))
        elif kind == "note":
            story.append(Paragraph(escape(blk[1]), S["note"]))
        elif kind == "bullets":
            story.append(ListFlowable([ListItem(Paragraph(rich(b), S["body"]), leftIndent=12)
                                       for b in blk[1]], bulletType="bullet", start="•", leftIndent=12))
        elif kind == "mono":
            story.append(Paragraph("<br/>".join(escape(x) for x in blk[1]), S["mono"]))
        elif kind == "kv":
            data = [[Paragraph(escape(k), S["cellb"]),
                     Paragraph(escape(v), S["mono"] if mono else S["cell"])] for k, v, mono in blk[1]]
            t = Table(data, colWidths=[1.45 * inch, width - 1.45 * inch])
            t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.4, H["rule"]),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]))
            story.append(t)
        elif kind == "table":
            headers, rows = blk[1], blk[2]
            if headers[0] == "#":
                fr = [0.05, 0.21, 0.15, 0.08, 0.51]
            elif headers[0] == "Cycle":
                fr = [0.1, 0.16, 0.26, 0.48]
            else:
                fr = [1 / len(headers)] * len(headers)
            data = [[Paragraph(escape(h), S["head"]) for h in headers]]
            for r in rows:
                data.append([Paragraph(rich(v) if headers[j] == "Cause" else escape(v), S["cell"])
                             for j, v in enumerate(r)])
            t = Table(data, colWidths=[f * width for f in fr], repeatRows=1)
            style = [("BACKGROUND", (0, 0), (-1, 0), H["head_bg"]),
                     ("LINEBELOW", (0, 0), (-1, -1), 0.4, H["rule"]),
                     ("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
            style += [("BACKGROUND", (0, i), (-1, i), H["band"]) for i in range(2, len(data), 2)]
            t.setStyle(TableStyle(style))
            story.append(t)
            story.append(Spacer(1, 4))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(H["muted"])
        canvas.drawString(0.8 * inch, 0.5 * inch, f"{APP_NAME} v{VERSION}")
        canvas.drawRightString(letter[0] - 0.8 * inch, 0.5 * inch, f"Page {doc.page}")
        canvas.setStrokeColor(H["rule"])
        canvas.line(0.8 * inch, 0.65 * inch, letter[0] - 0.8 * inch, 0.65 * inch)
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                            topMargin=0.7 * inch, bottomMargin=0.85 * inch,
                            title="SEL Trip Event Summary", author=APP_NAME)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def export_report(text, path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        export_docx(text, path)
    elif ext == ".pdf":
        export_pdf(text, path)
    else:
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)


# --------------------------------------------------------------------------
# GUI
# --------------------------------------------------------------------------
PALETTES = {
    "light": dict(bg="#fafafa", surface="#ffffff", text="#1b1b1b", muted="#616161",
                  border="#d9d9d9", accent="#005fb8", drop="#f4f8fd", drop_hover="#e3eefb",
                  drop_line="#9fb9d9", sel="#cce4f7", trip="#c42b1c", ok="#0f7b0f", warn="#9a5800"),
    "dark": dict(bg="#1c1c1c", surface="#252525", text="#e9e9e9", muted="#9e9e9e",
                 border="#3a3a3a", accent="#4cc2ff", drop="#20262d", drop_hover="#22344a",
                 drop_line="#4d6a8b", sel="#264f78", trip="#ff8f85", ok="#6ccb5f", warn="#ffb74d"),
}

PLACEHOLDER = ("Notes to include at the top of the summary, e.g.\n"
               "Crew found a cracked insulator at pole 47. Recommend checking Z1 reach.")
HELP_TEXT = """QUICK START
1. Drag your relay's event files (or a whole folder) onto the drop zone.
   Best: the .CEV event report plus the relay's .RDB settings file.
2. Read the summary. Lines marked CHECK are things the program is unsure about.
3. Edit the text if needed, then Copy all or Export to Word / PDF.
New here? Click "Try samples" to load example files.

WHICH FILES TO USE
  .CEV        Compressed event report - the most complete (elements, currents, settings)
  .RDB        QuickSet settings database - adds the trip equation (TR) and line length (LL)
  .EVE .TXT   Text event report (terminal EVE command or a saved capture)
  .SUM .HDR   Event summary / COMTRADE header written by the relay
  .HIS        Event history list - one line per event, least detail
  .SER        Sequential events recorder - element on/off times
Getting a CEV: download the event from the relay with AcSELerator QuickSet and
choose the compressed (CEV) format, or save the output of the CEV command from a
terminal session.

EVENT TYPES (what the relay decided the fault was)
  AG BG CG       single phase-to-ground fault (A, B or C)
  AB BC CA       phase-to-phase fault
  ABG BCG CAG    two phases to ground
  ABC            three-phase fault
  TRIP           tripped but no fault type found (transfer trip, manual/SCADA, external input)
  ER             event recorded by the ER trigger - no trip
  TRIG           manually triggered event
  A trailing  T  (for example "AG T") means the relay tripped.

COMMON TARGETS (front-panel lights)
  INST instantaneous    TIME time-delayed    COMM communications-assisted
  ZONE1/2/3 distance zones   50 instantaneous overcurrent   51 time-overcurrent
  81 frequency    SOTF switch-onto-fault    A B C G N faulted phases / ground
  RS recloser reset (ready)    CY recloser cycling    LO recloser locked out

ORGANIZING THE SUMMARY
  Show      All events, trips only, needs review, faults only, or no-trip events
            (the numbers are how many events each choice includes)
  Relay     Only one relay's events
  Sort      Oldest first, newest first, by relay, or needs review first
  Detail    Summary only (overview + one line per event), Standard (adds the
            details and CHECKs), or Full (adds element timeline, settings, SER)
  Hide duplicates   Skip history rows already covered by a detailed report
  Search    Only events whose details contain the text (Ctrl+F)
  The summary states which view it shows. Export > Spreadsheet (.csv) gives one
  row per event with the same filter and order, ready to sort in Excel.

CONFIDENCE AND CHECKS
  High    element data + settings, and nothing looked wrong
  Medium  element data or targets, settings missing or one item to check
  Low     history row, very little data, or several items to check
  CHECK lines say exactly what to confirm: missing data (e.g. no settings),
  disagreements (targets vs. event type vs. currents), or unusual results
  (reverse or beyond-line location, slow breaker, off-nominal frequency).

FAULT LOCATION
  Shown in the relay's line-length units. With settings loaded it is also
  shown as % of the line. Negative = behind the relay. $$$$$ = not calculated.

SHORTCUTS
  Ctrl+O add files    F5 generate    Ctrl+Shift+C copy    Ctrl+S Word    Ctrl+P PDF
  Ctrl+F search    Ctrl+D dark mode    F1 this help    Delete removes the selected files

The program only reads your files. Nothing is uploaded, except feedback you choose
to send and the update check (which only asks GitHub for the latest version).
"""

EMPTY_HINT = ("Drop SEL event files or a folder above and the trip summary will appear here.\n"
              "New here? Click \"Try samples\" (under the file list) to load example files.\n\n"
              "Everything in this box is editable before you copy or save it.")


def settings_path():
    base = (os.environ.get("APPDATA") or os.environ.get("XDG_CONFIG_HOME")
            or os.path.join(os.path.expanduser("~"), ".config"))
    return os.path.join(base, "SEL Trip Analyzer", "settings.json")


def load_settings():
    import json
    try:
        with open(settings_path(), encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_settings(data):
    import json
    try:
        os.makedirs(os.path.dirname(settings_path()), exist_ok=True)
        with open(settings_path(), "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2)
    except Exception:
        pass


def system_prefers_dark():
    if sys.platform != "win32":
        return False
    try:
        import winreg
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                            r"Software\Microsoft\Windows\CurrentVersion\Themes\Personalize") as k:
            return winreg.QueryValueEx(k, "AppsUseLightTheme")[0] == 0
    except Exception:
        return False


def resource_path(*parts):
    """Locate bundled files both from source and inside the PyInstaller .exe."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, *parts)


def set_window_icon(root):
    try:
        ico = resource_path("assets", "icon.ico")
        if sys.platform == "win32" and os.path.exists(ico):
            root.iconbitmap(default=ico)
            return
        png = resource_path("assets", "icon.png")
        if os.path.exists(png):
            import tkinter as tk
            root._icon_img = tk.PhotoImage(file=png)
            root.iconphoto(True, root._icon_img)
    except Exception:
        pass


def set_windows_titlebar(root, dark):
    if sys.platform != "win32":
        return
    try:
        import ctypes
        root.update_idletasks()
        hwnd = ctypes.windll.user32.GetParent(root.winfo_id())
        val = ctypes.c_int(1 if dark else 0)
        for attr in (20, 19):  # DWMWA_USE_IMMERSIVE_DARK_MODE (new, old builds)
            if ctypes.windll.dwmapi.DwmSetWindowAttribute(
                    hwnd, attr, ctypes.byref(val), ctypes.sizeof(val)) == 0:
                break
    except Exception:
        pass


def run_gui():
    import threading
    import tkinter as tk
    import tkinter.font as tkfont
    from tkinter import ttk, filedialog, messagebox
    from types import SimpleNamespace

    if sys.platform == "win32":  # crisp text on high-DPI screens
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass

    try:
        from tkinterdnd2 import TkinterDnD, DND_FILES
        root = TkinterDnD.Tk()
        dnd_ok = True
    except Exception:
        root = tk.Tk()
        dnd_ok = False
    try:
        import sv_ttk
    except Exception:
        sv_ttk = None

    cfg = load_settings()
    st = {"dark": cfg.get("dark", system_prefers_dark()), "update": None, "hover": False,
          "placeholder": True, "empty": True, "pal": None}

    root.title(APP_NAME)
    set_window_icon(root)
    if sys.platform == "win32":  # own taskbar group/icon instead of python's
        try:
            import ctypes
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("SELTripAnalyzer.App")
        except Exception:
            pass
    root.geometry(cfg.get("geometry", "1180x920"))
    root.minsize(920, 720)

    # ---- fonts: Segoe UI Variable on Windows 11, sensible fallbacks elsewhere
    fams = set(tkfont.families(root))

    def pick(*names):
        return next((n for n in names if n in fams), tkfont.nametofont("TkDefaultFont").actual("family"))

    UI = pick("Segoe UI Variable Text", "Segoe UI", "Helvetica Neue", "DejaVu Sans", "Inter")
    UI_SEMI = pick("Segoe UI Variable Display Semibold", "Segoe UI Semibold", "Inter Display SemiBold", UI)
    MONO = pick("Cascadia Mono", "Cascadia Code", "Consolas", "JetBrains Mono", "DejaVu Sans Mono")
    F = {
        "body": (UI, 10), "small": (UI, 9), "cap": (UI, 9, "bold"),
        "title": (UI_SEMI, 17, "bold" if UI_SEMI == UI else "normal"),
        "drop": (UI_SEMI, 13, "bold" if UI_SEMI == UI else "normal"),
        "mono": (MONO, 10), "mono_b": (MONO, 10, "bold"), "mono_h": (MONO, 12, "bold"),
    }

    style = ttk.Style(root)

    # ================= layout =================
    root.columnconfigure(0, weight=1)
    root.rowconfigure(4, weight=1)

    # ---- header
    hdr = ttk.Frame(root, padding=(22, 16, 22, 10))
    hdr.grid(row=0, column=0, sticky="ew")
    hdr.columnconfigure(0, weight=1)
    ttk.Label(hdr, text=APP_NAME, style="Title.TLabel").grid(row=0, column=0, sticky="w")
    ttk.Label(hdr, text="Trip-cause summaries from SEL relay event reports and QuickSet settings",
              style="Sub.TLabel").grid(row=1, column=0, sticky="w", pady=(2, 0))
    v_dark = tk.BooleanVar(value=st["dark"])
    dark_sw = ttk.Checkbutton(hdr, text="Dark mode", variable=v_dark,
                              style="Switch.TCheckbutton" if sv_ttk else "TCheckbutton")
    help_btn = ttk.Button(hdr, text="Help")
    help_btn.grid(row=0, column=1, rowspan=2, padx=(0, 14))
    dark_sw.grid(row=0, column=2, rowspan=2, padx=(0, 18))
    upd_btn = ttk.Button(hdr, text="Check for updates")
    upd_btn.grid(row=0, column=3, rowspan=2)

    # ---- drop zone
    drop = tk.Canvas(root, height=122, highlightthickness=0, bd=0, cursor="hand2")
    drop.grid(row=1, column=0, sticky="ew", padx=22, pady=(4, 0))

    # ---- files + comments
    mid = ttk.Frame(root, padding=(22, 16, 22, 0))
    mid.grid(row=2, column=0, sticky="ew")
    mid.columnconfigure(0, weight=1, uniform="mid")
    mid.columnconfigure(1, weight=1, uniform="mid")

    left = ttk.Frame(mid)
    left.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
    left.columnconfigure(0, weight=1)
    ttk.Label(left, text="FILES", style="Cap.TLabel").grid(row=0, column=0, sticky="w")
    count_lbl = ttk.Label(left, text="none loaded", style="Hint.TLabel")
    count_lbl.grid(row=0, column=1, sticky="e")
    lb = tk.Listbox(left, height=6, selectmode="extended", activestyle="none", relief="flat",
                    bd=0, highlightthickness=1, font=F["body"])
    lb.grid(row=1, column=0, columnspan=2, sticky="nsew", pady=(6, 8))
    fb = ttk.Frame(left)
    fb.grid(row=2, column=0, columnspan=2, sticky="w")

    right = ttk.Frame(mid)
    right.grid(row=0, column=1, sticky="nsew", padx=(10, 0))
    right.columnconfigure(0, weight=1)
    right.rowconfigure(1, weight=1)
    ttk.Label(right, text="COMMENTS & SUGGESTIONS", style="Cap.TLabel").grid(row=0, column=0, sticky="w")
    ttk.Label(right, text="added to the top of the summary", style="Hint.TLabel").grid(row=0, column=1, sticky="e")
    comments = tk.Text(right, height=6, wrap="word", relief="flat", bd=0, highlightthickness=1,
                       padx=10, pady=8, font=F["body"], undo=True)
    comments.grid(row=1, column=0, columnspan=2, sticky="nsew", pady=(6, 0))

    # ---- action row
    act = ttk.Frame(root, padding=(22, 14, 22, 8))
    act.grid(row=3, column=0, sticky="ew")
    view = {"show": cfg.get("show", "all"), "sort": cfg.get("sort", "time"),
            "detail": cfg.get("detail", "standard"), "relay": ""}
    for key, options in (("show", SHOW_OPTIONS), ("sort", SORT_OPTIONS), ("detail", DETAIL_OPTIONS)):
        if view[key] not in dict(options):
            view[key] = options[0][0] if key != "detail" else "standard"

    def combo(text, width):
        ttk.Label(act, text=text, style="Hint.TLabel").pack(side="left", padx=(0, 6))
        cb = ttk.Combobox(act, state="readonly", width=width)
        cb.pack(side="left", padx=(0, 14))
        return cb

    show_cb = combo("Show", 16)
    relay_cb = combo("Relay", 14)
    sort_cb = combo("Sort", 16)
    detail_cb = combo("Detail", 11)
    v_dups = tk.BooleanVar(value=cfg.get("hide_dups", True))
    dups_cb = ttk.Checkbutton(act, text="Hide duplicates", variable=v_dups)
    dups_cb.pack(side="left")
    gen_btn = ttk.Button(act, text="Generate summary", style="Accent.TButton")
    gen_btn.pack(side="right")

    # ---- output
    outf = ttk.Frame(root, padding=(22, 4, 22, 0))
    outf.grid(row=4, column=0, sticky="nsew")
    outf.columnconfigure(0, weight=1)
    outf.rowconfigure(1, weight=1)
    oh = ttk.Frame(outf)
    oh.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 6))
    ttk.Label(oh, text="TRIP SUMMARY", style="Cap.TLabel").pack(side="left")
    ttk.Label(oh, text="  editable", style="Hint.TLabel").pack(side="left")
    save_btn = ttk.Menubutton(oh, text="Export")
    save_btn.pack(side="right")
    export_menu = tk.Menu(save_btn, tearoff=False)
    save_btn["menu"] = export_menu
    copy_btn = ttk.Button(oh, text="Copy all")
    copy_btn.pack(side="right", padx=(0, 8))
    search_var = tk.StringVar()
    search_entry = ttk.Entry(oh, textvariable=search_var, width=26)
    search_entry.pack(side="right", padx=(0, 16))
    ttk.Label(oh, text="Search", style="Hint.TLabel").pack(side="right", padx=(0, 6))
    out = tk.Text(outf, wrap="word", undo=True, relief="flat", bd=0, highlightthickness=1,
                  padx=18, pady=14, font=F["mono"], spacing1=1, spacing3=1)
    out.grid(row=1, column=0, sticky="nsew")
    osb = ttk.Scrollbar(outf, orient="vertical", command=out.yview)
    osb.grid(row=1, column=1, sticky="ns", padx=(4, 0))
    out.configure(yscrollcommand=osb.set)

    # ---- feedback on the output
    fbk = ttk.Frame(root, padding=(22, 14, 22, 0))
    fbk.grid(row=5, column=0, sticky="ew")
    fbk.columnconfigure(0, weight=1)
    fh = ttk.Frame(fbk)
    fh.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 6))
    ttk.Label(fh, text="FEEDBACK ON THIS OUTPUT", style="Cap.TLabel").pack(side="left")
    ttk.Label(fh, text="  what was wrong, missing, or should be added?", style="Hint.TLabel").pack(side="left")
    feedback = tk.Text(fbk, height=3, wrap="word", relief="flat", bd=0, highlightthickness=1,
                       padx=10, pady=8, font=F["body"], undo=True)
    feedback.grid(row=1, column=0, sticky="ew")
    fside = ttk.Frame(fbk)
    fside.grid(row=1, column=1, sticky="ns", padx=(12, 0))
    v_attach = tk.BooleanVar(value=False)
    ttk.Checkbutton(fside, text="Attach summary", variable=v_attach).pack(anchor="w")
    send_btn = ttk.Button(fside, text="Send feedback")
    send_btn.pack(fill="x", pady=(8, 0))

    # ---- status bar
    sbar = ttk.Frame(root, padding=(22, 8, 22, 12))
    sbar.grid(row=6, column=0, sticky="ew")
    status = ttk.Label(sbar, text="Ready" + ("" if dnd_ok else "  -  drag-and-drop unavailable, use Add files"),
                       style="Hint.TLabel")
    status.pack(side="left")
    ttk.Label(sbar, text=f"v{VERSION}", style="Hint.TLabel").pack(side="right")

    # ================= theme =================
    def round_rect(c, x1, y1, x2, y2, r, **kw):
        pts = [x1 + r, y1, x2 - r, y1, x2, y1, x2, y1 + r, x2, y2 - r, x2, y2,
               x2 - r, y2, x1 + r, y2, x1, y2, x1, y2 - r, x1, y1 + r, x1, y1]
        return c.create_polygon(pts, smooth=True, **kw)

    def draw_drop(_=None):
        pal = st["pal"]
        drop.delete("all")
        drop.configure(bg=pal["bg"])
        w, h = max(drop.winfo_width(), 200), int(drop["height"])
        hov = st["hover"]
        round_rect(drop, 2, 2, w - 2, h - 2, 16, fill=pal["drop_hover"] if hov else pal["drop"],
                   outline=pal["accent"] if hov else pal["drop_line"], dash=(7, 5), width=1.6)
        cx, ic = w // 2, pal["accent"]
        drop.create_line(cx, 18, cx, 42, fill=ic, width=2.4, arrow="last", arrowshape=(9, 10, 5))
        drop.create_line(cx - 15, 36, cx - 15, 47, cx + 15, 47, cx + 15, 36, fill=ic, width=2.4,
                         capstyle="round", joinstyle="round")
        drop.create_text(cx, 72, text="Drop SEL event files or folders here", font=F["drop"], fill=pal["text"])
        drop.create_text(cx, 98, font=F["small"], fill=pal["muted"],
                         text="CEV  ·  EVE  ·  TXT  ·  SUM  ·  HIS  ·  SER  ·  HDR  ·  RDB        or click to browse")

    def style_text(w, pal):
        w.configure(highlightthickness=1, relief="flat", bd=0, bg=pal["surface"], fg=pal["text"], insertbackground=pal["text"],
                    selectbackground=pal["sel"], selectforeground=pal["text"],
                    highlightbackground=pal["border"], highlightcolor=pal["accent"])

    def apply_theme():
        dark = st["dark"]
        pal = st["pal"] = PALETTES["dark" if dark else "light"]
        if sv_ttk:
            sv_ttk.set_theme("dark" if dark else "light")
            if "Segoe UI Variable Text" not in fams:  # off Windows 11: map theme fonts
                for name in ("SunValleyCaptionFont", "SunValleyBodyFont", "SunValleyBodyStrongFont",
                             "SunValleyBodyLargeFont", "SunValleySubtitleFont", "SunValleyTitleFont"):
                    try:
                        tkfont.nametofont(name).configure(family=UI)
                    except Exception:
                        pass
        else:
            style.theme_use("clam")
            style.configure(".", background=pal["bg"], foreground=pal["text"], font=F["body"],
                            bordercolor=pal["border"], focuscolor=pal["accent"])
            style.configure("TButton", background=pal["surface"], padding=(12, 5), relief="flat")
            style.map("TButton", background=[("active", pal["drop_hover"])])
            style.configure("Accent.TButton", background=pal["accent"],
                            foreground="#000000" if dark else "#ffffff")
            style.map("Accent.TButton", background=[("active", pal["accent"])])
            style.configure("TCheckbutton", background=pal["bg"], foreground=pal["text"],
                            indicatorbackground=pal["surface"], indicatorforeground=pal["accent"])
            style.map("TCheckbutton", background=[("active", pal["bg"])])
            style.configure("TScrollbar", background=pal["surface"], troughcolor=pal["bg"],
                            arrowcolor=pal["muted"])
        style.configure("Title.TLabel", font=F["title"], foreground=pal["text"])
        style.configure("Sub.TLabel", font=F["body"], foreground=pal["muted"])
        style.configure("Cap.TLabel", font=F["cap"], foreground=pal["muted"])
        style.configure("Hint.TLabel", font=F["small"], foreground=pal["muted"])
        root.configure(bg=pal["bg"])
        restyle_widgets()
        root.after(60, restyle_widgets)  # the theme re-colours plain Tk widgets on <<ThemeChanged>>
        draw_drop()
        set_windows_titlebar(root, dark)

    def restyle_widgets():
        pal = st["pal"]
        comments.configure(font=F["body"])
        out.configure(font=F["mono"])
        lb.configure(font=F["body"], highlightthickness=1, relief="flat", bd=0)
        style_text(comments, pal)
        style_text(out, pal)
        style_text(feedback, pal)
        feedback.configure(font=F["body"])
        feedback.tag_configure("ph", foreground=pal["muted"])
        lb.configure(bg=pal["surface"], fg=pal["text"], selectbackground=pal["sel"],
                     selectforeground=pal["text"], highlightbackground=pal["border"],
                     highlightcolor=pal["accent"])
        comments.tag_configure("ph", foreground=pal["muted"])
        out.tag_configure("h1", font=F["mono_h"], foreground=pal["text"])
        out.tag_configure("h2", font=F["mono_b"], foreground=pal["accent"])
        out.tag_configure("rule", foreground=pal["border"])
        out.tag_configure("label", foreground=pal["muted"])
        out.tag_configure("trip", font=F["mono_b"], foreground=pal["trip"])
        out.tag_configure("notrip", font=F["mono_b"], foreground=pal["ok"])
        out.tag_configure("warn", foreground=pal["warn"])
        out.tag_configure("warn_b", font=F["mono_b"], foreground=pal["warn"])
        out.tag_configure("bold", font=F["mono_b"])
        out.tag_configure("hint", font=F["body"], foreground=pal["muted"])

    # ================= behaviour =================
    files = []

    def set_status(text):
        status.config(text=text)

    def refresh_list():
        lb.delete(0, "end")
        for p in files:
            lb.insert("end", f"  {os.path.basename(p)}")
        count_lbl.config(text=f"{len(files)} loaded" if files else "none loaded")

    def add_paths(paths):
        known = {os.path.normcase(os.path.abspath(f)) for f in files}
        new = [p for p in expand_paths(paths) if os.path.normcase(os.path.abspath(p)) not in known]
        files.extend(new)
        refresh_list()
        if new:
            generate()
        elif paths:
            set_status("No new SEL files found in what was dropped.")

    def browse(_=None):
        sel = filedialog.askopenfilenames(
            title="Select SEL event / RDB files",
            filetypes=[("SEL event & settings", " ".join("*" + e for e in sorted(ALL_EXTS))),
                       ("All files", "*.*")])
        if sel:
            add_paths(list(sel))

    def browse_folder():
        d = filedialog.askdirectory(title="Select a folder of SEL files")
        if d:
            add_paths([d])

    def remove_sel():
        for i in reversed(lb.curselection()):
            files.pop(i)
        refresh_list()
        if files:
            generate()

    def clear_all():
        files.clear()
        st["cache"] = None
        view["relay"] = ""
        refresh_choices(None)
        refresh_list()
        show_empty()
        set_status("Cleared")

    def get_comments():
        return "" if st["placeholder"] else comments.get("1.0", "end-1c")

    def show_placeholder():
        if not comments.get("1.0", "end-1c").strip():
            comments.delete("1.0", "end")
            comments.insert("1.0", PLACEHOLDER, "ph")
            st["placeholder"] = True

    def hide_placeholder(_=None):
        if st["placeholder"]:
            comments.delete("1.0", "end")
            st["placeholder"] = False

    def show_empty():
        out.delete("1.0", "end")
        out.insert("1.0", EMPTY_HINT, "hint")
        st["empty"] = True

    def highlight():
        for tag in ("h1", "h2", "rule", "label", "trip", "notrip", "bold", "warn", "warn_b"):
            out.tag_remove(tag, "1.0", "end")
        for i, ln in enumerate(out.get("1.0", "end-1c").split("\n"), 1):
            s = ln.strip()
            if not s:
                continue
            if i == 1 and s.startswith("SEL TRIP"):
                out.tag_add("h1", f"{i}.0", f"{i}.end")
            elif re.fullmatch(r"[=\-]{10,}", s):
                out.tag_add("rule", f"{i}.0", f"{i}.end")
            elif ln == s and re.fullmatch(r"[A-Z][A-Z /&()\-]{3,}", s):
                out.tag_add("h2", f"{i}.0", f"{i}.end")
            elif re.match(r"Event \d+ of \d+", ln):
                out.tag_add("bold", f"{i}.0", f"{i}.end")
            elif s.startswith("Probable cause"):
                out.tag_add("bold", f"{i}.0", f"{i}.end")
            else:
                m = re.match(r"^(\s{2}[A-Za-z][A-Za-z ./]{2,12}:)", ln)
                if m:
                    out.tag_add("label", f"{i}.0", f"{i}.{len(m.group(1))}")
            k = ln.find("CHECK:")
            if k >= 0:
                out.tag_add("warn", f"{i}.{k}", f"{i}.end")
                out.tag_add("warn_b", f"{i}.{k}", f"{i}.{k + 6}")
            m = re.search(r"\[\d+ to check\]", ln)
            if m:
                out.tag_add("warn_b", f"{i}.{m.start()}", f"{i}.{m.end()}")
            if s.startswith(("Needs review:", "Confidence")) and "to check" in s:
                out.tag_add("warn", f"{i}.0", f"{i}.end")
            for word, tag in (("TRIP - ", "trip"), ("Relay TRIPPED", "trip"),
                              ("No trip - ", "notrip"), ("No trip recorded", "notrip")):
                k = ln.find(word)
                if k >= 0:
                    out.tag_add(tag, f"{i}.{k}", f"{i}.{k + len(word.rstrip(' -'))}")

    # ---- view options (filter / sort / detail)
    def current_opts():
        return {"show": view["show"], "relay": view["relay"], "sort": view["sort"], "detail": view["detail"],
                "hide_duplicates": v_dups.get(), "search": search_var.get().strip()}

    def loaded(reload=False):
        """Parse the files once; view changes reuse the result."""
        key = tuple(files)
        if reload or not st.get("cache") or st["cache"][0] != key:
            events, rdb, errors = load_files(files)
            st["cache"] = (key, events, rdb, errors, view_choices(events, rdb))
        return st["cache"][1:]

    def refresh_choices(choices=None):
        counts = choices["counts"] if choices else {}
        show_cb["values"] = [f"{lbl} ({counts[k]})" if choices else lbl for k, lbl in SHOW_OPTIONS]
        show_cb.current([k for k, _ in SHOW_OPTIONS].index(view["show"]))
        relays = choices["relays"] if choices else []
        st["relay_keys"] = [""] + [k for k, _ in relays]
        relay_cb["values"] = ([f"All relays ({sum(n for _, n in relays)})"] if relays else ["All relays"]) + \
                             [f"{k} ({n})" for k, n in relays]
        if view["relay"] not in st["relay_keys"]:
            view["relay"] = ""
        relay_cb.current(st["relay_keys"].index(view["relay"]))
        sort_cb["values"] = [lbl for _, lbl in SORT_OPTIONS]
        sort_cb.current([k for k, _ in SORT_OPTIONS].index(view["sort"]))
        detail_cb["values"] = [lbl for _, lbl in DETAIL_OPTIONS]
        detail_cb.current([k for k, _ in DETAIL_OPTIONS].index(view["detail"]))

    def edits_ok():
        """Changing the view rebuilds the summary; check before throwing away hand edits."""
        if st["empty"] or not out.edit_modified():
            return True
        return messagebox.askyesno(
            "Discard your edits?",
            "Changing the view rebuilds the summary, so the edits you made in the summary box will be lost.\n\n"
            "Tip: copy or export your edited summary first.\n\nRebuild the summary now?")

    def on_view(kind):
        keys = {"show": [k for k, _ in SHOW_OPTIONS], "sort": [k for k, _ in SORT_OPTIONS],
                "detail": [k for k, _ in DETAIL_OPTIONS], "relay": st.get("relay_keys", [""])}[kind]
        cb = {"show": show_cb, "sort": sort_cb, "detail": detail_cb, "relay": relay_cb}[kind]
        new = keys[cb.current()] if cb.current() >= 0 else keys[0]
        if new == view[kind]:
            return
        if not edits_ok():
            cb.current(keys.index(view[kind]))
            return
        view[kind] = new
        if files:
            generate(reload=False)

    def on_dups():
        if not edits_ok():
            v_dups.set(not v_dups.get())
            return
        if files:
            generate(reload=False)

    def on_search(_=None):
        if st.get("search_job"):
            root.after_cancel(st["search_job"])
        st["search_job"] = root.after(350, run_search)

    def run_search():
        st["search_job"] = None
        if search_var.get().strip() == st.get("last_search", ""):
            return
        if not edits_ok():
            search_var.set(st.get("last_search", ""))
            return
        if files:
            generate(reload=False)

    def generate(_=None, reload=True):
        if not files:
            set_status("Add some files first")
            return
        root.config(cursor="watch")
        root.update_idletasks()
        try:
            events, rdb, errors, choices = loaded(reload)
            opts = current_opts()
            st["last_search"] = opts["search"]
            refresh_choices(choices)
            txt = build_report(events, rdb, errors, files, get_comments(), opts)
            out.delete("1.0", "end")
            out.insert("1.0", txt)
            out.edit_modified(False)
            st["empty"] = False
            highlight()
            n = len([e for e in events if e.kind != "ser"])
            _, shown, _ = prepare_view(events, rdb, opts)
            showing = f"showing {len(shown)} of {n}" if len(shown) != n else f"{n} event{'s' if n != 1 else ''}"
            set_status(f"{showing} from {len(files)} file{'s' if len(files) != 1 else ''}"
                       + (f"  -  {len(errors)} could not be read" if errors else ""))
        finally:
            root.config(cursor="")

    def flash(btn, text, original):
        btn.config(text=text)
        root.after(1600, lambda: btn.config(text=original))

    def copy_all(_=None):
        if st["empty"]:
            return
        root.clipboard_clear()
        root.clipboard_append(out.get("1.0", "end-1c"))
        flash(copy_btn, "Copied ✓", "Copy all")
        set_status("Summary copied to clipboard")

    EXPORT_TYPES = {".docx": ("Word document", "*.docx"), ".pdf": ("PDF", "*.pdf"),
                    ".csv": ("Spreadsheet (CSV)", "*.csv"), ".txt": ("Text file", "*.txt")}

    def export_as(ext):
        if st["empty"]:
            set_status("Generate a summary first")
            return
        name, pattern = EXPORT_TYPES[ext]
        stamp = datetime.datetime.now().strftime("%Y-%m-%d")
        p = filedialog.asksaveasfilename(
            title=f"Export summary as {name}", defaultextension=ext,
            initialfile=f"Trip summary {stamp}{ext}",
            initialdir=cfg.get("export_dir") or os.path.expanduser("~"),
            filetypes=[(name, pattern), ("All files", "*.*")])
        if not p:
            return
        cfg["export_dir"] = os.path.dirname(p)
        root.config(cursor="watch")
        root.update_idletasks()
        try:
            if ext == ".csv":   # built from the data (same filter and order as the summary)
                events, rdb, _, _ = loaded()
                with open(p, "w", encoding="utf-8-sig", newline="") as f:
                    f.write(report_csv(events, rdb, current_opts()))
            else:
                export_report(out.get("1.0", "end-1c"), p)
        except PermissionError:
            messagebox.showerror("Export failed", f"Couldn't write {os.path.basename(p)}.\n"
                                                  "Is it open in Word or a PDF viewer? Close it and try again.")
            return
        except Exception as e:
            messagebox.showerror("Export failed", f"{type(e).__name__}: {e}")
            return
        finally:
            root.config(cursor="")
        flash(save_btn, "Exported ✓", "Export")
        set_status(f"Saved {p}" + ("  (one row per event; edits in the summary box aren't included)" if ext == ".csv" else ""))
        if ext != ".txt" and sys.platform == "win32":
            try:
                os.startfile(p)  # open it in Word / the PDF viewer
            except Exception:
                pass

    def toggle_dark(_=None):
        if _ is not None:  # keyboard shortcut
            v_dark.set(not v_dark.get())
        st["dark"] = v_dark.get()
        apply_theme()

    # ---- updater
    def run_bg(work, done):
        def worker():
            try:
                res, err = work(), None
            except Exception as e:
                res, err = None, e
            root.after(0, lambda: done(res, err))
        threading.Thread(target=worker, daemon=True).start()

    def show_update_available(info):
        st["update"] = info
        upd_btn.config(text=f"Update to v{info['version']}", style="Accent.TButton")
        set_status(f"Version {info['version']} is available")

    def check_updates(silent=False):
        if "OWNER/" in GITHUB_REPO:
            if not silent:
                messagebox.showinfo("Updates", "No GitHub repository is configured in this build.")
            return
        upd_btn.config(state="disabled")
        if not silent:
            set_status("Checking GitHub for updates...")

        def done(info, err):
            upd_btn.config(state="normal")
            if err:
                if not silent:
                    set_status("Update check failed")
                    messagebox.showerror("Updates", f"Could not reach GitHub:\n{err}")
                return
            if not info:
                if not silent:
                    set_status(f"You're on the latest version (v{VERSION})")
                    messagebox.showinfo("Updates", f"You're up to date (v{VERSION}).")
                return
            show_update_available(info)
            if not silent:
                prompt_update(info)

        run_bg(check_for_update, done)

    def prompt_update(info):
        notes = info["notes"][:1200] or "(no release notes)"
        if messagebox.askyesno(
                "Update available",
                f"Version {info['version']} is available (you have {VERSION}).\n\n"
                f"What's new:\n{notes}\n\nDownload and install it now?\n"
                "The program will restart when finished."):
            do_install(info)

    def do_install(info):
        upd_btn.config(state="disabled")

        def progress(done_b, total):
            msg = (f"Downloading update... {100 * done_b // total}%" if total
                   else f"Downloading update... {done_b // 1024} KB")
            root.after(0, lambda: set_status(msg))

        def done(finish, err):
            upd_btn.config(state="normal")
            if err:
                set_status("Update failed")
                messagebox.showerror("Update failed", str(err))
                return
            set_status("Update installed - restarting...")
            on_close(save_only=True)
            finish()
            root.destroy()

        run_bg(lambda: install_update(info, progress), done)

    def on_update_click():
        if st["update"]:
            prompt_update(st["update"])
        else:
            check_updates(silent=False)

    FEEDBACK_PH = ("e.g. Event 3 was a DTT from the remote end, not a local fault. "
                   "Would be nice to show the breaker number in the quick summary.")
    st["fb_ph"] = False

    def fb_placeholder(show):
        if show and not feedback.get("1.0", "end-1c").strip():
            feedback.delete("1.0", "end")
            feedback.insert("1.0", FEEDBACK_PH, "ph")
            st["fb_ph"] = True
        elif not show and st["fb_ph"]:
            feedback.delete("1.0", "end")
            st["fb_ph"] = False

    def send_feedback():
        import urllib.parse
        import webbrowser
        text = "" if st["fb_ph"] else feedback.get("1.0", "end-1c").strip()
        if not text:
            set_status("Type your feedback first")
            feedback.focus_set()
            return
        summary = "" if st["empty"] else out.get("1.0", "end-1c")
        stamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
        exts = sorted({os.path.splitext(p)[1].lower() for p in files})
        body = (f"{text}\n\n---\nApp version: v{VERSION}  \n"
                f"Files loaded: {len(files)} ({', '.join(exts) or 'none'})\n")
        if v_attach.get() and summary:
            room = 6000 - len(urllib.parse.quote(body))
            clip = summary
            while clip and len(urllib.parse.quote(clip)) > room:
                clip = clip[: int(len(clip) * 0.8)]
            body += ("\n<details><summary>Summary output</summary>\n\n```\n" + clip
                     + ("\n... (truncated)" if len(clip) < len(summary) else "") + "\n```\n</details>\n")
        title = "Feedback: " + re.sub(r"\s+", " ", text)[:70]
        if v_attach.get() and summary and "OWNER/" not in GITHUB_REPO and not messagebox.askyesno(
                "Attach summary?",
                "Feedback is posted as a public GitHub issue, so anyone can read the attached "
                "summary (relay names, dates, locations).\n\nAttach it anyway?"):
            return
        # keep a local copy of everything that's sent
        log = os.path.join(os.path.dirname(settings_path()), "feedback.log")
        try:
            os.makedirs(os.path.dirname(log), exist_ok=True)
            with open(log, "a", encoding="utf-8") as f:
                f.write(f"===== {stamp}  v{VERSION}\n{text}\n\n")
        except Exception:
            pass
        if "OWNER/" in GITHUB_REPO:
            messagebox.showinfo("Feedback", f"Saved to {log}")
            return
        url = (f"https://github.com/{GITHUB_REPO}/issues/new?"
               + urllib.parse.urlencode({"title": title, "body": body}))
        webbrowser.open(url)
        feedback.delete("1.0", "end")
        st["fb_ph"] = False
        fb_placeholder(True)
        flash(send_btn, "Opened ✓", "Send feedback")
        set_status("Feedback opened on GitHub - click 'Create' there to submit (a copy was saved locally)")

    def load_samples():
        folder = resource_path("sample_files")
        if not os.path.isdir(folder):
            messagebox.showinfo("Samples", "Sample files aren't included in this build.")
            return
        add_paths([folder])
        set_status("Loaded the sample files - these are made-up events for trying the program")

    def show_help(_=None):
        if st.get("help_win") and st["help_win"].winfo_exists():
            st["help_win"].lift()
            return
        pal = st["pal"]
        win = tk.Toplevel(root)
        st["help_win"] = win
        win.title(f"{APP_NAME} - Help")
        win.geometry("760x700")
        win.configure(bg=pal["bg"])
        set_window_icon(win)
        frm = ttk.Frame(win, padding=16)
        frm.pack(fill="both", expand=True)
        frm.rowconfigure(0, weight=1)
        frm.columnconfigure(0, weight=1)
        t = tk.Text(frm, wrap="word", relief="flat", bd=0, highlightthickness=1, padx=16, pady=12,
                    font=F["mono"])
        t.grid(row=0, column=0, sticky="nsew")
        sb2 = ttk.Scrollbar(frm, orient="vertical", command=t.yview)
        sb2.grid(row=0, column=1, sticky="ns", padx=(4, 0))
        t.configure(yscrollcommand=sb2.set)
        style_text(t, pal)
        t.tag_configure("h", font=F["mono_b"], foreground=pal["accent"])
        for line in HELP_TEXT.splitlines():
            tag = "h" if line and line == line.upper() and line[0].isalpha() else ""
            t.insert("end", line + "\n", tag)
        t.configure(state="disabled")
        bar2 = ttk.Frame(frm)
        bar2.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(10, 0))
        ttk.Button(bar2, text="Online guide", command=lambda: __import__("webbrowser").open(
            f"https://github.com/{GITHUB_REPO}#readme")).pack(side="left")
        ttk.Button(bar2, text="Close", command=win.destroy).pack(side="right")
        win.bind("<Escape>", lambda e: win.destroy())

    def on_close(save_only=False):
        save_settings({"dark": st["dark"], "geometry": root.geometry(), "export_dir": cfg.get("export_dir"),
                       "show": view["show"], "sort": view["sort"], "detail": view["detail"],
                       "hide_dups": v_dups.get()})
        if not save_only:
            root.destroy()

    # ---- wiring
    upd_btn.config(command=on_update_click)
    help_btn.config(command=show_help)
    dark_sw.config(command=toggle_dark)
    gen_btn.config(command=generate)
    copy_btn.config(command=copy_all)
    for ext, label in ((".docx", "Word document (.docx)"), (".pdf", "PDF (.pdf)"),
                       (".csv", "Spreadsheet for Excel (.csv)"), (".txt", "Plain text (.txt)")):
        export_menu.add_command(label=label, command=lambda e=ext: export_as(e))
    send_btn.config(command=send_feedback)
    for kind, cb in (("show", show_cb), ("relay", relay_cb), ("sort", sort_cb), ("detail", detail_cb)):
        cb.bind("<<ComboboxSelected>>", lambda e, k=kind: (on_view(k), e.widget.selection_clear()))
    dups_cb.config(command=on_dups)
    search_entry.bind("<KeyRelease>", on_search)
    search_entry.bind("<Escape>", lambda e: (search_var.set(""), on_search()))
    root.bind("<Control-f>", lambda e: (search_entry.focus_set(), search_entry.select_range(0, "end"), "break")[-1])
    feedback.bind("<FocusIn>", lambda e: fb_placeholder(False))
    feedback.bind("<FocusOut>", lambda e: fb_placeholder(True))
    feedback.bind("<Control-a>", lambda e: (e.widget.tag_add("sel", "1.0", "end"), "break")[1])
    for txt, cmd in (("Add files", browse), ("Add folder", browse_folder),
                     ("Remove", remove_sel), ("Clear", clear_all), ("Try samples", load_samples)):
        ttk.Button(fb, text=txt, command=cmd).pack(side="left", padx=(0, 6))

    drop.bind("<Configure>", draw_drop)
    drop.bind("<Button-1>", browse)
    drop.bind("<Enter>", lambda e: (st.update(hover=True), draw_drop()))
    drop.bind("<Leave>", lambda e: (st.update(hover=False), draw_drop()))
    comments.bind("<FocusIn>", hide_placeholder)
    comments.bind("<FocusOut>", lambda e: show_placeholder())
    lb.bind("<Delete>", lambda e: remove_sel())

    if dnd_ok:
        def on_drop(event):
            st["hover"] = False
            draw_drop()
            add_paths(list(root.tk.splitlist(event.data)))
            return event.action

        def on_enter(event):
            st["hover"] = True
            draw_drop()
            return event.action

        def on_leave(event):
            st["hover"] = False
            draw_drop()
            return event.action

        for w in (drop, lb, out, root):
            w.drop_target_register(DND_FILES)
            w.dnd_bind("<<Drop>>", on_drop)
        drop.dnd_bind("<<DropEnter>>", on_enter)
        drop.dnd_bind("<<DropLeave>>", on_leave)

    def select_all(e):
        e.widget.tag_add("sel", "1.0", "end")
        return "break"

    out.bind("<Control-a>", select_all)
    comments.bind("<Control-a>", select_all)
    root.bind("<Control-o>", browse)
    root.bind("<Control-Return>", generate)
    root.bind("<F5>", generate)
    root.bind("<Control-d>", toggle_dark)
    root.bind("<F1>", show_help)
    root.bind("<Control-s>", lambda e: export_as(".docx"))
    root.bind("<Control-p>", lambda e: export_as(".pdf"))
    root.bind("<Control-Shift-C>", copy_all)
    root.protocol("WM_DELETE_WINDOW", on_close)

    apply_theme()
    refresh_choices(None)
    fb_placeholder(True)
    show_placeholder()
    show_empty()
    refresh_list()

    # hooks used by scripts/screenshots.py
    root.app = SimpleNamespace(add_paths=add_paths, comments=comments, hide_placeholder=hide_placeholder,
                               feedback=feedback, fb_placeholder=fb_placeholder,
                               generate=generate, toggle_dark=toggle_dark, v_dark=v_dark,
                               show_update_available=show_update_available)

    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if args:  # files dropped onto the .exe icon
        root.after(150, lambda: add_paths(args))
    if os.environ.get("SEL_TRIP_NO_UPDATE_CHECK") != "1":
        root.after(1500, lambda: check_updates(silent=True))
    root.mainloop()


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--cli":
        # --cli <files/folders> [--comments TEXT] [--out FILE.docx|.pdf|.txt|.csv ...]
        #       [--show all|trips|review|faults|notrip] [--relay NAME] [--sort time|newest|relay|review]
        #       [--detail summary|standard|full] [--search TEXT] [--hide-duplicates] [--format text|csv]
        args, outs, note, opts, fmt = sys.argv[2:], [], "", {}, "text"
        while "--out" in args:
            i = args.index("--out")
            outs.append(args[i + 1])
            del args[i:i + 2]
        for flag, key in (("--comments", None), ("--show", "show"), ("--relay", "relay"), ("--sort", "sort"),
                          ("--detail", "detail"), ("--search", "search"), ("--format", "format")):
            if flag in args:
                i = args.index(flag)
                if key is None:
                    note = args[i + 1]
                elif key == "format":
                    fmt = args[i + 1]
                else:
                    opts[key] = args[i + 1]
                del args[i:i + 2]
        if "--hide-duplicates" in args:
            args.remove("--hide-duplicates")
            opts["hide_duplicates"] = True
        paths = expand_paths(args)
        ev, rdb, err = load_files(paths)
        report = build_report(ev, rdb, err, paths, note, opts)
        for o in outs:
            if o.lower().endswith(".csv"):
                with open(o, "w", encoding="utf-8-sig", newline="") as f:
                    f.write(report_csv(ev, rdb, opts))
            else:
                export_report(report, o)
        if sys.stdout is not None:  # the windowed .exe has no console
            print(report_csv(ev, rdb, opts) if fmt == "csv" else report, end="" if fmt == "csv" else "\n")
    else:
        run_gui()
